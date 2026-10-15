"""Run: python tests/e2e.py. Requires installed Chromium, Playwright and python-docx."""
from pathlib import Path
import io, os, tempfile, zipfile, re, json
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from playwright.sync_api import sync_playwright
ROOT=Path(__file__).resolve().parent.parent
def page_scripts(name):
    return re.findall(r'<script\s+src="([^"]+)"\s*></script>',(ROOT/name).read_text(encoding='utf-8'))

STORAGE_SHIM="""(init)=>{
  // set_content pages have an opaque origin without localStorage; this keeps
  // the app's own storage code testable (dump_storage carries it to the next page).
  try{localStorage.getItem('x');return;}catch{}
  const data=new Map(Object.entries(init||{}));
  const api={getItem:k=>data.has(String(k))?data.get(String(k)):null,setItem:(k,v)=>{data.set(String(k),String(v));},removeItem:k=>{data.delete(String(k));},
    clear:()=>data.clear(),key:i=>[...data.keys()][i]??null,get length(){return data.size;},dump:()=>Object.fromEntries(data)};
  Object.defineProperty(window,'localStorage',{value:api,configurable:true});
}"""

def dump_storage(page):
    return page.evaluate("()=>localStorage.dump?localStorage.dump():Object.fromEntries(Object.keys(localStorage).map(k=>[k,localStorage.getItem(k)]))")

def load_page(page,name,scripts=None,storage=None):
    # Scripts are taken from the page itself so the test never drifts from
    # the shipped HTML (an explicit list is still accepted).
    scripts=scripts or page_scripts(name)
    # Browser administration policy blocks URL navigation in this environment.
    # The actual extension HTML and JS are nevertheless executed in Chromium
    # after removing only external script tags and CSP (opaque test origin).
    html=(ROOT/name).read_text(encoding='utf-8')
    html=re.sub(r'<meta[^>]+http-equiv="Content-Security-Policy"[^>]*>', '', html)
    html=re.sub(r'<script\s+src="[^"]+"\s*></script>', '', html)
    html=re.sub(r'<link\s+rel="stylesheet"[^>]*>', '', html)
    page.set_content(html,wait_until='load')
    page.evaluate(STORAGE_SHIM,storage or {})
    # Stylesheets as the page lists them (themes.css, workspace.css, reverse.css…).
    for css_name in re.findall(r'<link rel="stylesheet" href="([^"]+)"',(ROOT/name).read_text(encoding='utf-8')):
        page.add_style_tag(content=(ROOT/css_name).read_text(encoding='utf-8'))
    for script in scripts:page.add_script_tag(content=(ROOT/script).read_text(encoding='utf-8'))

def write_source(p):
    d=Document()
    for sec in d.sections:
        sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.4)
        sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.2)
    d.add_paragraph('T.C.')
    h=d.add_paragraph()
    r=h.add_run('ŞANLIURFA 1. ASLİYE HUKUK MAHKEMESİNE');r.underline=True;r.bold=False
    d.add_paragraph('DAVACI:')
    d.add_paragraph('Ali Veli')
    d.add_paragraph('DAVALI: Karşı taraf')
    d.add_paragraph('VEKİLİ: Av. Fatih Nakışlı')
    d.add_paragraph('AÇIKLAMALAR')
    d.add_paragraph('I. TEMEL AÇIKLAMALAR')
    v=d.add_paragraph();v.alignment=WD_ALIGN_PARAGRAPH.CENTER
    v.paragraph_format.line_spacing=1.5
    v.paragraph_format.space_before=Pt(13)
    v.add_run('HMK ');r=v.add_run('m. 119');r.bold=True;v.add_run(' hükmü uygulanmalıdır.')
    cite=d.add_paragraph();cite.add_run('Yargıtay 3.');cite.add_run(' Hukuk Dairesi 2025/5112 E., ');cite.add_run('2026/3190 K.')
    d.add_paragraph('SONUÇ VE İSTEM')
    d.add_paragraph('1. Davanın KABULÜNE karar verilmesini talep ederim.')
    d.add_paragraph('Av. Fatih Nakışlı')
    d.add_paragraph('EKLER')
    d.add_paragraph('EK-1 Örnek belge')
    d.save(p)

def write_petition(p):
    """Dumanbey-type petition: separate label lines, Word auto-numbered requests."""
    d=Document()
    for t in ['T.C.','İSTANBUL 22. ASLİYE TİCARET MAHKEMESİNE','İHTİYATİ HACİZ KARARINA İTİRAZ VE İVEDİ İNCELEME TALEPLİDİR.',
              'DOSYA NO: 2026/713 E.','KARAR NO: 2026/708','İHTİYATİ HACZE İTİRAZ EDEN:',
              'Dumanbey Rezistans Isı Sistemleri ve Sanayi Ticaret Limited Şirketi','ADRES: Başpınar OSB, Şehitkamil/Gaziantep',
              'VEKİLİ: Av. Ayşe Nur Geyik Öztürkmen','ADRES VE UETS: 16615-16448-62422',
              'İHTİYATİ HACİZ TALEP EDEN: Pimar Plastik Kimya Tekstil Makina Sanayi ve Ticaret Limited Şirketi',
              'VEKİLİ: Av. Emin Kablama','İLGİLİ İCRA DOSYASI: İstanbul Anadolu 10. İcra Dairesi, 2026/18304 E.',
              'KONU: İhtiyati haciz kararına itiraz edilmesi istemidir.','AÇIKLAMALAR:','Mahkemenizce verilen karara itiraz ediyoruz.',
              'I. MAHKEMENİN YETKİSİNE İTİRAZ EDİLMEKTEDİR.','Müvekkil şirketin merkezi Gaziantep’tir.',
              'HUKUKİ SEBEPLER: İİK m. 50, 257, 258, 259 ve 265; HMK m. 6.','HUKUKİ DELİLLER: Çek, imza sirküleri, ticaret sicili kayıtları.',
              'SONUÇ VE İSTEM:','Yukarıda arz edilen sebeplerle;']:
        d.add_paragraph(t)
    for t in ['Yetki itirazının kabulüne,','İhtiyati haciz kararının kaldırılmasına,','Yargılama giderlerinin karşı tarafa yükletilmesine,']:
        d.add_paragraph(t,style='List Number')
    for t in ['Karar verilmesini vekâleten talep ederim.','14.10.2025','İtiraz Eden Vekili','Av. Ayşe Nur Geyik Öztürkmen','EKLER:','1- İhtiyati haciz kararı örneği.','2- Çek örneği.']:
        d.add_paragraph(t)
    d.save(p)

def data_xml(p):
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None
        return etree.fromstring(z.read('content.xml'))

def one(el,tag):
    return el.xpath(f'./{tag}')[0]

def udf_text(root):
    return ''.join(one(root,'content').itertext())

def fetch_data(page,source):
    page.locator('#fileInput').set_input_files(str(source))
    page.wait_for_timeout(1200)
    print('UPLOAD STATUS',page.locator('#status').text_content())
    print('WARNINGS',page.locator('#warningList').all_text_contents())
    print('READY',page.locator('#download').is_enabled())
    if page.locator('#review').count() and not page.locator('#reviewLabel').is_hidden():page.locator('#review').check()
    page.wait_for_function("!document.getElementById('download').disabled",timeout=7000)

def save_download(page,button,path):
    with page.expect_download(timeout=20000) as info: page.locator(button).click()
    dl=info.value
    dl.save_as(str(path))
    assert path.exists() and path.stat().st_size>100
    return dl.suggested_filename

def main():
    with tempfile.TemporaryDirectory(prefix='wordtoudf-e2e-') as tmp, sync_playwright() as pw:
        tmp=Path(tmp);source=tmp/'sample.docx';write_source(source)
        browser=pw.chromium.launch(executable_path=os.environ.get('CHROMIUM','/opt/pw-browsers/chromium' if os.path.exists('/opt/pw-browsers/chromium') else '/usr/bin/chromium'),headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
        context=browser.new_context(accept_downloads=True,viewport={'width':1500,'height':1000})
        page=context.new_page();errors=[]
        page.on('pageerror',lambda e:errors.append(str(e)))
        load_page(page,'workspace.html')
        fetch_data(page,source)
        assert page.locator('[data-role="court-heading"]').count()==1
        assert page.locator('[data-role="signature"]').count()==1
        assert page.locator('[data-role="section-heading"]').count()==1
        assert page.locator('[data-role="metadata"]').count()>=3
        assert not page.locator('[data-role="metadata"]').all_text_contents()[0].count('DAVALI')
        legal_xml=tmp/'legal.udf'
        fname=save_download(page,'#download',legal_xml)
        assert fname=='sample.udf',fname
        root=data_xml(legal_xml)
        full=udf_text(root)
        assert 'HMK m. 119' in full and '2026/3190 K.' in full
        assert 'DAVACI\t: Ali VELİ' in full
        assert 'VEKİLİ\t: Av. Fatih NAKIŞLI (UETS: 16615-16448-62422)' in full
        assert b'Cambria' not in etree.tostring(root) and all(c.get('family') in (None,'Times New Roman') for c in root.iter('content'))
        assert full.count('ŞANLIURFA')==1
        paras=root.xpath('./elements/paragraph')
        court=[p for p in paras if any('ŞANLIURFA' in full[int(r.get('startOffset')):int(r.get('startOffset'))+int(r.get('length'))] for r in p)]
        assert len(court)==1 and court[0].get('Alignment')=='1'
        assert all(r.get('underline')=='false' for r in court[0] if r.tag=='content')
        assert abs(float(one(one(root,'properties'),'pageFormat').get('leftMargin'))-42.52)<.02
        print('PASS hukuk profili DOCX→UDF, metin, roller, kenar boşlukları, UDF ZIP/XML')
        # Avoid changing the paragraph's source formatting in faithful mode.
        page.locator('#typographyProfile').select_option('faithful')
        page.wait_for_function("document.getElementById('sectionPanel').hidden")
        assert page.locator('[data-role="court-heading"]').count()==0
        faithful=tmp/'faithful.udf';save_download(page,'#download',faithful)
        root2=data_xml(faithful)
        props=one(one(root2,'properties'),'pageFormat')
        assert abs(float(props.get('leftMargin'))-(2.3*72/2.54))<.03
        paras2=root2.xpath('./elements/paragraph')
        styles=[p for p in paras2 if p.get('Alignment')=='1' and p.get('SpaceAbove')=='13']
        assert len(styles)>=1,'Kaynak Word paragraf aralığı/hizalama korunmadı'
        assert 'DAVACI:\nAli Veli' in udf_text(root2)
        print('PASS Word biçimini koru: paragraf hizası/aralığı, kenar boşlukları')
        # Re-enter legal, upload complete DOCX with warning and verify explicit acknowledgement.
        page.locator('#typographyProfile').select_option('legal')
        assert page.locator('#download').is_enabled()
        # UDF→DOCX reverse converter end-to-end.
        reverse=context.new_page();reverse.on('pageerror',lambda e:errors.append(str(e)))
        load_page(reverse,'reverse.html',['zip.js','udf-reader.js','docx-writer.js','reverse.js'])
        reverse.locator('#fileInput').set_input_files(str(legal_xml))
        reverse.wait_for_function("!document.getElementById('downloadWord').disabled",timeout=20000)
        round_trip=tmp/'roundtrip.docx'
        with reverse.expect_download(timeout=20000) as info:reverse.locator('#downloadWord').click()
        info.value.save_as(str(round_trip))
        back=Document(round_trip)
        entire='\n'.join(p.text for p in back.paragraphs)
        assert 'ŞANLIURFA' in entire and 'Yargıtay 3. Hukuk Dairesi' in entire
        assert 'Av. Fatih NAKIŞLI' in entire
        print('PASS UDF→DOCX geriye dönüş, temel metin bütünlüğü')
        # Negative ZIP archive case: traversal entry must be rejected.
        traversal=tmp/'path-traversal.docx'
        with zipfile.ZipFile(source) as zin,zipfile.ZipFile(traversal,'w') as zout:
            for name in zin.namelist():zout.writestr(name,zin.read(name))
            zout.writestr('../disallowed.xml','malicious')
        page.locator('#fileInput').set_input_files(str(traversal))
        page.wait_for_function("document.querySelector('#status').classList.contains('error')",timeout=7000)
        assert 'geçersiz dosya yolu' in page.locator('#status').text_content()
        assert page.locator('#download').is_disabled()
        print('PASS kötü niyetli ZIP ../ yolu reddedildi')
        # XML entity/doctype must be rejected, including otherwise valid DOCX ZIP.
        dtd=tmp/'doctype.docx'
        with zipfile.ZipFile(source) as zin,zipfile.ZipFile(dtd,'w') as zout:
            for name in zin.namelist():
                data=zin.read(name)
                if name=='word/document.xml':
                    data=data.replace(b'<w:document',b'<!DOCTYPE w:document [<!ENTITY secret "test">]><w:document',1)
                zout.writestr(name,data)
        page.locator('#fileInput').set_input_files(str(dtd))
        page.wait_for_function("document.querySelector('#status').classList.contains('error')",timeout=7000)
        assert 'DTD' in page.locator('#status').text_content() or 'ENTITY' in page.locator('#status').text_content()
        assert page.locator('#download').is_disabled()
        print('PASS DOCX XML DTD/ENTITY reddedildi')
        # Simulate an untrusted UDF whose compressed payload fails CRC/inflate.
        corrupt=tmp/'corrupt.udf'
        blob=bytearray(legal_xml.read_bytes())
        blob[len(blob)//2]^=0x31
        corrupt.write_bytes(blob)
        reverse.locator('#fileInput').set_input_files(str(corrupt))
        reverse.wait_for_function("document.querySelector('#status').classList.contains('error')",timeout=7000)
        assert reverse.locator('#downloadWord').is_disabled()
        print('PASS bozuk UDF ZIP reddedildi')
        # A larger legal petition should finish synchronously without truncation.
        large=tmp/'large.docx'
        big=Document();big.add_paragraph('AÇIKLAMALAR')
        for i in range(1500):big.add_paragraph(f'Paragraf {i+1}: Deliller usulüne uygun değerlendirilmelidir.')
        big.add_paragraph('SONUÇ VE İSTEM');big.add_paragraph('1. Kabulüne karar verilmesini talep ederim.')
        big.save(large)
        page.locator('#fileInput').set_input_files(str(large))
        page.wait_for_function("document.getElementById('previewCount').textContent.includes('1503 paragraf')",timeout=25000)
        if not page.locator('#reviewLabel').is_hidden():page.locator('#review').check()
        output=tmp/'large.udf';save_download(page,'#download',output)
        root_large=data_xml(output)
        data=udf_text(root_large)
        assert 'Paragraf 1:' in data and 'Paragraf 1500:' in data
        # v1.4.0: "1. Kabulüne karar verilmesini talep ederim." becomes a list item
        # plus a separate dated closing paragraph, so 1503 source paragraphs → 1504.
        assert len(root_large.xpath('./elements/paragraph'))==1504
        print('PASS 1500 gövde paragrafı: 1504 UDF paragrafı (kapanış ayrıldı), metin kaybı yok')
        # v1.4.0 dilekçe standardı: ortak sekme, tek DOSYA NO, UETS, satır içi bölümler, gerçek listeler, tarih.
        petition=tmp/'petition.docx';write_petition(petition)
        page.locator('#listMode').select_option('auto')  # bu bölüm gerçek UDF listelerini denetler
        page.locator('#fileInput').set_input_files(str(petition))
        page.wait_for_function("document.getElementById('status').textContent.includes('petition.docx')",timeout=7000)
        if not page.locator('#reviewLabel').is_hidden():page.locator('#review').check()
        assert page.locator('#changeList li').count()>=10
        xs=page.evaluate("""()=>[...document.querySelectorAll('.paper .paragraph')].map(p=>{const t=p.querySelector('.tab-stop');const c=t&&t.nextElementSibling;return c&&c.textContent===':'?c.getBoundingClientRect().left:null}).filter(x=>x!==null)""")
        assert len(xs)==13 and max(xs)-min(xs)<0.5,xs
        if os.environ.get('SHOT'):page.locator('#paper').screenshot(path=os.environ['SHOT'])
        markers=page.locator('.list-marker').all_text_contents()
        assert markers==['1.','2.','3.','1.','2.'],markers
        out=tmp/'petition.udf';save_download(page,'#download',out)
        save_dir=os.environ.get('SAVE_DIR')
        if save_dir:
            import shutil
            shutil.copy(petition,Path(save_dir)/'ornek-dilekce-kaynak.docx');shutil.copy(out,Path(save_dir)/'ornek-dilekce.udf')
            save_download(page,'#reportDownload',Path(save_dir)/'ornek-dilekce-donusum-raporu.txt')
        root_p=data_xml(out);txt=udf_text(root_p)
        paras=root_p.xpath('./elements/paragraph')
        def ptext(p):return ''.join(txt[int(r.get('startOffset')):int(r.get('startOffset'))+int(r.get('length'))] for r in p)
        lines=[ptext(p).rstrip('\n') for p in paras]
        for expected in ['DOSYA NO\t: 2026/713 E. 2026/708 K.','İTİRAZ EDEN\t: Dumanbey Rezistans Isı Sistemleri ve Sanayi Ticaret Limited Şirketi',
                         'ADRESİ\t: Başpınar OSB, Şehitkamil/Gaziantep','VEKİLİ\t: Av. Ayşe Nur GEYİK ÖZTÜRKMEN (UETS: 16615-16448-62422)',
                         'TALEP EDEN\t: Pimar Plastik Kimya Tekstil Makina Sanayi ve Ticaret Limited Şirketi','VEKİLİ\t: Av. Emin KABLAMA',
                         'İCRA DOSYA NO\t: İstanbul Anadolu 10. İcra Dairesi, 2026/18304 E.','KONU\t: İhtiyati haciz kararına itiraz edilmesi istemidir.',
                         'AÇIKLAMALAR\t:','HUKUKİ SEBEPLER\t: İİK m. 50, 257, 258, 259 ve 265; HMK m. 6.','HUKUKİ DELİLLER\t: Çek, imza sirküleri, ticaret sicili kayıtları.',
                         'SONUÇ VE İSTEM\t: Yukarıda arz edilen sebeplerle;','EKLER\t:']:
            assert expected in lines,(expected,lines)
        assert 'KARAR NO' not in txt and 'ADRES VE UETS' not in txt
        field=[p for p,l in zip(paras,lines) if '\t:' in l]
        assert len(field)==13
        stops={p.get('TabSet') for p in field};assert len(stops)==1,stops
        assert all(p.get('LeftIndent')=='0' and p.get('FirstLineIndent')=='0' for p in field)
        # UYAP'ın kendi yazımı: "ETİKET<TAB>:" tek, kalın ve altı çizili <content>; ayrı <tab> öğesi yok.
        assert not root_p.xpath('.//tab'),'ayrı <tab> öğesi yazılmaz'
        for p in field:
            kids=list(p);seg=lambda k:txt[int(k.get('startOffset')):int(k.get('startOffset'))+int(k.get('length'))]
            assert kids[0].tag=='content' and seg(kids[0]).rstrip('\n').endswith('\t:') and seg(kids[0]).count('\t')==1,seg(kids[0])
            assert kids[0].get('underline')=='true' and kids[0].get('bold')=='true'
            for k in kids[1:]:assert k.get('underline')=='false'
        lists=[(p.get('ListId'),l) for p,l in zip(paras,lines) if p.get('Numbered')=='true']
        assert lists==[('1','Yetki itirazının kabulüne,'),('1','İhtiyati haciz kararının kaldırılmasına,'),('1','Yargılama giderlerinin karşı tarafa yükletilmesine,'),
                       ('2','İhtiyati haciz kararı örneği.'),('2','Çek örneği.')],lists
        import datetime,zoneinfo
        today=datetime.datetime.now(zoneinfo.ZoneInfo('Europe/Istanbul')).strftime('%d.%m.%Y')
        assert f'Karar verilmesini vekâleten talep ederim. {today}' in lines and '14.10.2025' not in txt
        sig=[p for p,l in zip(paras,lines) if l in ('İtiraz Eden Vekili','Av. Ayşe Nur GEYİK ÖZTÜRKMEN')]
        assert len(sig)==2 and all(p.get('Alignment')=='2' for p in sig)
        body=[p for p,l in zip(paras,lines) if l.startswith('Mahkemenizce')][0]
        assert body.get('LineSpacing')=='0.5' and body.get('Alignment')=='3'
        print('PASS dilekçe standardı: ortak sekme, tek DOSYA NO, UETS, satır içi bölümler, otomatik listeler, tarih, imza')
        # v1.4.1 varsayılan: metin numarası + sekme (numara ilk sözcüğün üstüne binmez).
        page.locator('#listMode').select_option('text')
        out_t=tmp/'petition-text.udf';save_download(page,'#download',out_t)
        root_t=data_xml(out_t);txt_t=udf_text(root_t)
        paras_t=root_t.xpath('./elements/paragraph')
        lines_t=[''.join(txt_t[int(r.get('startOffset')):int(r.get('startOffset'))+int(r.get('length'))] for r in p).rstrip('\n') for p in paras_t]
        assert not any(p.get('Numbered') for p in paras_t)
        items_t=[(p,l) for p,l in zip(paras_t,lines_t) if re.match(r'^\d+\.\t',l)]
        assert [l for _,l in items_t]==['1.\tYetki itirazının kabulüne,','2.\tİhtiyati haciz kararının kaldırılmasına,','3.\tYargılama giderlerinin karşı tarafa yükletilmesine,','1.\tİhtiyati haciz kararı örneği.','2.\tÇek örneği.'],items_t
        assert all(p.get('LeftIndent')=='18' and p.get('FirstLineIndent')=='-18' and p.get('TabSet')=='18:0:0' and txt_t[int(p[0].get('startOffset')):].split('\t')[0].endswith('.') for p,_ in items_t)
        assert not root_t.xpath('.//tab'),'numaradan sonraki sekme de metnin içinde'
        print('PASS metin numaralı liste: 1.<TAB> asılı girinti, Numbered özniteliği yok')
        page.locator('#listMode').select_option('auto')
        # Word → UDF → Word: lists stay real Word lists with separate numbering.
        reverse.locator('#fileInput').set_input_files(str(out))
        reverse.wait_for_function("!document.getElementById('downloadWord').disabled",timeout=20000)
        rt=tmp/'petition-roundtrip.docx'
        with reverse.expect_download(timeout=20000) as info:reverse.locator('#downloadWord').click()
        info.value.save_as(str(rt))
        if save_dir:shutil.copy(rt,Path(save_dir)/'ornek-dilekce-udf-den-word.docx')
        with zipfile.ZipFile(rt) as z:
            assert 'word/numbering.xml' in z.namelist()
            docxml=etree.fromstring(z.read('word/document.xml'))
        W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        numbered=[]
        for p in docxml.iter(W+'p'):
            num=p.find(f'./{W}pPr/{W}numPr')
            if num is not None:numbered.append((num.find(W+'numId').get(W+'val'),num.find(W+'ilvl').get(W+'val'),''.join(t.text or '' for t in p.iter(W+'t'))))
        assert [n[2] for n in numbered]==[l for _,l in lists],numbered
        assert len({n[0] for n in numbered[:3]})==1 and len({n[0] for n in numbered[3:]})==1 and numbered[0][0]!=numbered[3][0]
        assert all(n[1]=='0' for n in numbered)
        print('PASS Word→UDF→Word: talepler ve ekler iki ayrı gerçek Word listesi, düzey 0')
        assert not errors,errors
        print('PASS tarayıcı JS konsolu: hata yok')
        v15_checks(browser,tmp,petition,source)
        v16_paste(browser,tmp)
        v17_no_jump(browser,petition)
        v18_marks(browser,tmp,petition)
        v19_library(browser,tmp)
        v110_tools(browser,tmp,petition)
        v1101_audit(browser,tmp,petition)
        v111_design(browser,tmp,petition)
        v111_editing(browser,tmp,petition)
        v112_features(browser,tmp,petition)
        v112_share(browser,petition)
        csp_check(browser)
        browser.close()

def udf_lines(path):
    root=data_xml(path);txt=udf_text(root)
    return root,[''.join(txt[int(r.get('startOffset')):int(r.get('startOffset'))+int(r.get('length'))] for r in p).rstrip('\n') for p in root.xpath('./elements/paragraph')]

def v15_checks(browser,tmp,petition,sample):
    import datetime,zoneinfo
    today=datetime.datetime.now(zoneinfo.ZoneInfo('Europe/Istanbul')).strftime('%d.%m.%Y')
    # --- Telefon: 390 px, karanlık mod; kağıt beyaz ve ekrana sığıyor.
    ctx=browser.new_context(accept_downloads=True,viewport={'width':390,'height':844},color_scheme='dark')
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html')
    fetch_data(page,petition)
    sw=page.evaluate("()=>[document.getElementById('previewViewport').scrollWidth,document.getElementById('previewViewport').clientWidth,document.documentElement.scrollWidth,innerWidth]")
    assert sw[0]<=sw[1]+1 and sw[2]<=sw[3]+1,sw
    bg=page.evaluate("()=>getComputedStyle(document.getElementById('paper')).backgroundColor")
    assert bg in ('rgb(255, 255, 255)','rgba(255, 255, 255, 1)'),bg
    assert 'Sığdır' in page.locator('#zoomToggle').text_content()
    page.locator('#zoomToggle').click()
    ov=page.evaluate("()=>{const v=document.getElementById('previewViewport');return [getComputedStyle(v).overflowX,v.scrollWidth>v.clientWidth]}")
    assert ov[0]=='auto' and ov[1],ov  # gerçek boyutta yana kaydırılabilir
    page.locator('#zoomToggle').click()
    print('PASS telefon görünümü: önizleme ekrana sığıyor, kağıt karanlık modda da beyaz')
    # --- Denetim paneli görünür ve hata sayısı yazılı.
    assert page.locator('#checksPanel').is_visible() and page.locator('#checkList li').count()>=1
    assert 'hata' in page.locator('#checkStats').text_content()
    # --- Elle düzeltme: imzaya dokun → düz paragraf → UDF'de sola yaslı.
    page.locator('#paper .paragraph[data-role="signature"]').last.click()
    sheet=page.locator('.edit-sheet')
    assert sheet.is_visible()
    page.wait_for_timeout(350)  # açılış animasyonu (0,2 sn) bitsin
    box=sheet.bounding_box();assert box['y']+box['height']<=844+1 and box['width']>=380,box  # telefonda alttan açılan panel
    sheet.locator('.chip',has_text='Düz paragraf').click()
    assert page.locator('#paper .paragraph.edited').count()==1
    assert page.locator('#paper .paragraph.edited[data-role="edit-plain"]').count()==1
    page.locator('.edit-sheet .primary-btn',has_text='Tamam').click()
    assert page.locator('.edit-sheet').count()==0
    # Talep maddesi sil → numaralar kayar; geri al → geri gelir.
    items=page.locator('#paper .paragraph[data-role="result-item"]')
    assert items.count()==3
    items.nth(0).click();page.locator('.edit-sheet .chip',has_text='Paragrafı sil').click()
    nums=[re.match(r'^(\d+\.)',t).group(1) for t in page.locator('#paper .paragraph[data-role="result-item"]').all_inner_texts()]
    assert nums==['1.','2.'],nums
    assert page.locator('.edit-sheet').count()==0
    page.locator('#previewTools button',has_text='Geri al').click()
    assert page.locator('#paper .paragraph[data-role="result-item"]').count()==3
    # Çoklu seçim: iki gövde paragrafı ortala + kalın.
    page.locator('#previewTools button',has_text='Çoklu seçim').click()
    bodies=page.locator('#paper .paragraph[data-role="explanation-body"]')
    bodies.nth(0).click();bodies.nth(1).click()
    assert 'Seçilenleri düzenle' in page.locator('.select-bar').inner_text() and '2 paragraf' in page.locator('.select-bar').inner_text()
    page.locator('.select-bar button',has_text='Seçilenleri düzenle').click()
    page.locator('.edit-sheet .segmented .chip',has_text='Ortala').click()
    page.locator('.edit-sheet .chip',has_text='Kalın').click()
    al=page.evaluate("()=>[...document.querySelectorAll('#paper .paragraph.edited')].map(p=>p.style.textAlign)")
    assert al.count('center')==2,al
    page.locator('.edit-sheet .primary-btn',has_text='Tamam').click()
    page.locator('#previewTools button',has_text='Çoklu seçim').click()
    # Metin düzenleme (tek paragraf).
    page.locator('#paper .paragraph[data-role="topic"]').click()
    page.locator('.edit-sheet textarea').fill('KONU\t: İhtiyati haciz kararına itiraz ve ivedi inceleme istemidir.')
    page.locator('.edit-sheet button',has_text='Metni uygula').click()
    assert 'ivedi inceleme istemidir' in page.locator('#paper .paragraph.edited',has_text='KONU').inner_text()
    page.locator('.edit-sheet .primary-btn',has_text='Tamam').click()
    assert any('Elle düzenleme' in t for t in page.locator('#changeList li').all_text_contents())
    out=tmp/'override.udf';save_download(page,'#download',out)
    root,lines=udf_lines(out)
    paras=root.xpath('./elements/paragraph')
    sig=[p for p,l in zip(paras,lines) if l=='Av. Ayşe Nur GEYİK ÖZTÜRKMEN']
    assert len(sig)==1 and sig[0].get('Alignment')=='0',[(p.attrib,l) for p,l in zip(paras,lines)][-6:]
    assert [l for l in lines if re.match(r'^\d\.\t',l)][:3]==['1.\tYetki itirazının kabulüne,','2.\tİhtiyati haciz kararının kaldırılmasına,','3.\tYargılama giderlerinin karşı tarafa yükletilmesine,']
    centered=[p for p,l in zip(paras,lines) if p.get('Alignment')=='1' and l.startswith(('Mahkemenizce','Müvekkil'))]
    assert len(centered)==2 and all(c.get('bold')=='true' for p in centered for c in p if c.tag=='content'),[p.attrib for p in centered]
    assert 'KONU\t: İhtiyati haciz kararına itiraz ve ivedi inceleme istemidir.' in lines
    rp=tmp/'rapor.txt';save_download(page,'#reportDownload',rp);rt=rp.read_text(encoding='utf-8')
    assert 'Elle yapılan paragraf düzenlemeleri:' in rt and 'tür: Düz paragraf' in rt and 'metin değiştirildi' in rt,rt[-800:]
    print('PASS paragraf düzenleyici: tür, sil + numara kayması, geri al, çoklu seçim (ortala+kalın), metin; UDF ve raporda')
    # --- Düzenlenmiş DOCX indir.
    dx=tmp/'duzenlenmis.docx';save_download(page,'#downloadDocx',dx)
    d=Document(str(dx));texts=[p.text for p in d.paragraphs]
    assert 'VEKİLİ\t: Av. Ayşe Nur GEYİK ÖZTÜRKMEN (UETS: 16615-16448-62422)' in texts,texts[:12]
    assert any(t.startswith('Karar verilmesini vekâleten talep ederim. '+today) for t in texts)
    print('PASS düzenlenmiş Word (DOCX) indirme: UDF ile aynı metin')
    # --- Ayarlar dışa/içe aktarma.
    st=tmp/'ayarlar.json';save_download(page,'#exportSettings',st)
    import json
    data=json.loads(st.read_text(encoding='utf-8'));assert data['application']=='WordToUDF' and data['ownAttorney']['uets']=='16615-16448-62422'
    data['ownAttorney']={'name':'Av. Ayşe Nur GEYİK ÖZTÜRKMEN','uets':'99999-88888-77777'};st.write_text(json.dumps(data),encoding='utf-8')
    page.locator('#importSettings').set_input_files(str(st))
    page.wait_for_function("document.getElementById('ownUets').value==='99999-88888-77777'")
    assert page.locator('#paper .paragraph',has_text='UETS: 99999-88888-77777').count()==1
    print('PASS ayarlar dışa/içe aktarma: içe aktarılan vekil profili önizlemeye uygulandı')
    # --- Toplu dönüşüm: 2 geçerli + 1 hatalı dosya → ZIP + rapor.
    bad=tmp/'bozuk.docx';bad.write_bytes(b'PK'+b'0'*400)
    with page.expect_download(timeout=30000) as info:page.locator('#batchInput').set_input_files([str(petition),str(sample),str(bad)])
    zp=tmp/'toplu.zip';info.value.save_as(str(zp))
    with zipfile.ZipFile(zp) as z:
        names=sorted(z.namelist());assert names==['petition.udf','sample.udf','toplu-donusum-raporu.txt'],names
        rep=z.read('toplu-donusum-raporu.txt').decode('utf-8');assert '✘ bozuk.docx' in rep and '✔ petition.docx' in rep,rep
        inner=tmp/'p.udf';inner.write_bytes(z.read('petition.udf'));_,pl=udf_lines(inner);assert 'SONUÇ VE İSTEM\t: Yukarıda arz edilen sebeplerle;' in pl
    page.wait_for_function("document.getElementById('status').textContent.includes('2/3')")
    print('PASS toplu dönüşüm: 2/3 dosya UDF, hatalı dosya raporda, ZIP doğrulandı')
    assert not errors,errors
    ctx.close()
    # --- Şablon sayfası: mazeret dilekçesi.
    ctx=browser.new_context(accept_downloads=True,viewport={'width':1400,'height':1000})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'templates.html')
    page.locator('#court').fill('Şanlıurfa 8. Asliye Hukuk Mahkemesi');page.locator('#caseNo').fill('2026/1929')
    page.locator('#client').fill('İsa ÇİFTÇİ');page.locator('#other').fill('Yapı ve Kredi Bankası Anonim Şirketi')
    page.locator('#f_hearingDate').fill('2026-10-20');page.locator('#f_reason').select_option('sehir')
    page.wait_for_timeout(400)
    assert page.locator('#paper [data-role="court-heading"]').text_content()=="ŞANLIURFA 8. ASLİYE HUKUK MAHKEMESİ'NE"
    tu=tmp/'sablon.udf';save_download(page,'#download',tu)
    root,lines=udf_lines(tu)
    for exp in ['DOSYA NO\t: 2026/1929 Esas','DAVALI\t: İsa ÇİFTÇİ','VEKİLİ\t: Av. Fatih NAKIŞLI (UETS: 16615-16448-62422)','DAVACI\t: Yapı ve Kredi Bankası Anonim Şirketi',
                'KONU\t: 20.10.2026 tarihli duruşmaya ilişkin mesleki mazeretimizin sunulması ve kabulü talebidir.','SONUÇ VE İSTEM\t: Yukarıda arz edilen sebeplerle;',
                '1.\t20.10.2026 tarihli duruşmaya ilişkin mesleki mazeretimizin kabulüne,',f'Karar verilmesini saygıyla talep ederiz. {today}','Davalı İsa ÇİFTÇİ Vekili','Av. Fatih NAKIŞLI']:
        assert exp in lines,(exp,lines)
    assert not any('…' in l for l in lines),lines
    # Kullanıcı düzenlemesi korunur ve şablon değişince sıfırlanır.
    page.locator('#requests').fill('Mazeretimizin kabulüne,\nDuruşmanın ertelenmesine,')
    page.wait_for_timeout(300)
    assert page.locator('#paper [data-role="result-item"]').count()==2
    page.locator('#tpl').select_option('tedbir');page.wait_for_timeout(300)
    assert 'HMK m. 394' in page.locator('#body').input_value()
    td=tmp/'sablon.docx';save_download(page,'#downloadDocx',td)
    assert any('ihtiyati tedbir kararının kaldırılmasına' in p.text for p in Document(str(td)).paragraphs)
    # Bütün şablonlar önizleme üretir; şablon önizlemesinde de paragraf düzenleyici çalışır.
    ids=page.evaluate("()=>[...document.getElementById('tpl').options].map(o=>o.value)")
    assert len(ids)==14,ids
    for tid in ids:
        page.locator('#tpl').select_option(tid);page.wait_for_timeout(250)
        assert page.locator('#paper [data-role="court-heading"]').count()==1,tid
        assert page.locator('#paper [data-role="result-item"]').count()>=1,tid
        assert page.locator('#paper [data-role="result-closing"]').count()==1,tid
    page.locator('#paper .paragraph[data-role="result-closing"]').click()
    page.locator('.edit-sheet .segmented .chip',has_text='Sağa').click()
    assert page.locator('#paper .paragraph.edited').evaluate("p=>p.style.textAlign")=='right'
    page.keyboard.press('Escape')
    assert not errors,errors
    print('PASS şablon sayfası: 14 şablon, mazeret ve tedbir itirazı → UDF/DOCX, boş alan kalmadı, düzenleme korunuyor, paragraf düzenleyici çalışıyor')
    ctx.close()
    # --- Başlangıç sayfası ve PWA dosyaları.
    import json as _j
    m=_j.loads((ROOT/'app.webmanifest').read_text(encoding='utf-8'));assert m['start_url']=='workspace.html'
    swsrc=(ROOT/'sw.js').read_text(encoding='utf-8')
    for f in re.findall(r"'([\w.-]+\.(?:js|html|css|png|webmanifest))'",swsrc):assert (ROOT/f).exists(),f
    for html in ['workspace.html','templates.html','reverse.html','sure.html','index.html']:
        for src in page_scripts(html):assert src in swsrc,(html,src)
        for css in re.findall(r'<link\s+rel="stylesheet"\s+href="([^"]+)"',(ROOT/html).read_text(encoding='utf-8')):assert css in swsrc,(html,css)
    # Ana ekrana eklenebilmesi için 192 ve 512 px simgeler.
    sizes={i['sizes'] for i in m['icons']};assert {'192x192','512x512'}<=sizes,sizes
    for i in m['icons']:assert (ROOT/i['src']).exists() and i['src'] in swsrc,i
    print('PASS çevrimdışı web uygulaması: manifest ve önbellek listesi bütün dosyaları kapsıyor')

WORD_HTML="""<html xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:w="urn:schemas-microsoft-com:office:word"><head><style>p.MsoNormal{margin:0}</style></head><body lang=TR><!--StartFragment-->
<p class=MsoNormal align=center style='text-align:center'><b><span style='font-size:12.0pt'>ŞANLIURFA 8. ASLİYE HUKUK MAHKEMESİ'NE<o:p></o:p></span></b></p>
<p class=MsoNormal><b><u>DOSYA NO</u></b><span style='mso-tab-count:2'>                    </span><b><u>:</u></b> 2026/1929 Esas</p>
<p class=MsoNormal><b><u>DAVALI</u></b><span style='mso-tab-count:2'>       </span><b><u>:</u></b> İsa ÇİFTÇİ</p>
<p class=MsoNormal><b><u>VEKİLİ</u></b><span style='mso-tab-count:2'> </span><b><u>:</u></b> Av. Fatih NAKIŞLI</p>
<p class=MsoNormal><b><u>DAVACI</u></b><span style='mso-tab-count:2'> </span><b><u>:</u></b> Yapı ve Kredi Bankası Anonim Şirketi</p>
<p class=MsoNormal><b><u>KONU</u></b><span style='mso-tab-count:2'> </span><b><u>:</u></b> Duruşmaya ilişkin mazeretimizin sunulmasıdır.</p>
<p class=MsoNormal><b><u>AÇIKLAMALAR</u></b><span style='mso-tab-count:1'> </span><b><u>:</u></b></p>
<p class=MsoNormal style='text-align:justify;text-indent:35.4pt'>Mahkemenizin dosyasının 20.10.2026 tarihli duruşmasına <b>ivedi</b> bir mazeretimiz nedeniyle katılamayacağız.</p>
<p class=MsoNormal><b><u>SONUÇ VE İSTEM</u></b><span style='mso-tab-count:1'> </span><b><u>:</u></b> Yukarıda arz edilen sebeplerle;</p>
<p class=MsoListParagraphCxSpFirst style='mso-list:l0 level1 lfo1'><![if !supportLists]><span style='mso-list:Ignore'>1.<span style='font:7.0pt "Times New Roman"'>&nbsp;&nbsp;&nbsp;&nbsp;</span></span><![endif]>Mazeretimizin kabulüne,</p>
<p class=MsoListParagraphCxSpLast style='mso-list:l0 level1 lfo1'><![if !supportLists]><span style='mso-list:Ignore'>2.<span style='font:7.0pt "Times New Roman"'>&nbsp;&nbsp;&nbsp;</span></span><![endif]>Duruşmanın ertelenmesine,</p>
<p class=MsoNormal>Karar verilmesini saygıyla talep ederiz. 14.10.2026</p>
<p class=MsoNormal align=right style='text-align:right'><b>Davalı İsa Çiftçi Vekili</b></p>
<p class=MsoNormal align=right style='text-align:right'><b>Av. Fatih NAKIŞLI</b></p><img src="x.png">
<!--EndFragment--></body></html>"""
GDOCS_HTML='<meta charset="utf-8"><b style="font-weight:normal;" id="docs-internal-guid-1f2e"><p dir="ltr" style="text-align: center;"><span style="font-weight:700;">ANKARA 3. İŞ MAHKEMESİ\'NE</span></p><p dir="ltr"><span style="font-weight:700;text-decoration:underline;">DAVACI</span><span style="font-weight:400;">: Ali VELİ</span></p><p dir="ltr"><span style="font-weight:400;">Davacı ücret alacağı için </span><span style="font-weight:400;font-style:italic;">ayrıca</span><span style="font-weight:400;"> başvurmuştur.</span></p><p dir="ltr"><span>Yukarıda açıklanan nedenlerle davanın kabulüne karar verilmesini talep ederim.</span></p></b>'
PDF_TEXT="\n".join(['ŞANLIURFA 8. ASLİYE HUKUK MAHKEMESİ\'NE','DOSYA NO : 2026/1929 Esas','DAVALI : İsa ÇİFTÇİ','AÇIKLAMALAR :',
  'Mahkemenizin yukarıda esas numarası belirtilen dosyasının 14.10.2026 tarihli','duruşmasına, şehir dışında bulunmamız ve önceden planlanmış diğer mesleki iş ve','işlemlerimiz nedeniyle katılım sağlayamayacağımızdan mesleki mazeretimizin','kabulünü talep etmekteyiz.','Sayfa 1 / 1',
  'SONUÇ VE İSTEM : Yukarıda arz edilen sebeplerle;','1. Mazeretimizin kabulüne,','Karar verilmesini saygıyla talep ederiz.'])

def paste_into(page,html,text):
    page.evaluate("""([html,text])=>{const dt=new DataTransfer();if(html)dt.setData('text/html',html);dt.setData('text/plain',text);
      document.getElementById('pasteBox').dispatchEvent(new ClipboardEvent('paste',{clipboardData:dt,bubbles:true,cancelable:true}));}""",[html,text])
    page.wait_for_timeout(250)

def v16_paste(browser,tmp):
    import datetime,zoneinfo
    today=datetime.datetime.now(zoneinfo.ZoneInfo('Europe/Istanbul')).strftime('%d.%m.%Y')
    ctx=browser.new_context(accept_downloads=True,viewport={'width':1400,'height':1000})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html')
    # 1) Word'den biçimli kopya: sekme, liste numarası, kalın, hizalama, görsel.
    paste_into(page,WORD_HTML,'DOSYA NO\t: 2026/1929 Esas')
    assert 'Microsoft Word' in page.locator('#pasteSource').text_content()
    notes=' '.join(page.locator('#pasteNotes li').all_text_contents());assert 'görsel aktarılmadı' in notes,notes
    assert page.locator('#pasteBox').input_value().startswith('DOSYA NO')
    out=tmp/'word-paste.udf';save_download(page,'#download',out)
    root,lines=udf_lines(out)
    for e in ["ŞANLIURFA 8. ASLİYE HUKUK MAHKEMESİ'NE",'DOSYA NO\t: 2026/1929 Esas','DAVALI\t: İsa ÇİFTÇİ','VEKİLİ\t: Av. Fatih NAKIŞLI (UETS: 16615-16448-62422)',
              'KONU\t: Duruşmaya ilişkin mazeretimizin sunulmasıdır.','AÇIKLAMALAR\t:','1.\tMazeretimizin kabulüne,','2.\tDuruşmanın ertelenmesine,',f'Karar verilmesini saygıyla talep ederiz. {today}','Davalı İsa Çiftçi Vekili','Av. Fatih NAKIŞLI']:
        assert e in lines,(e,lines)
    txt=udf_text(root)
    ivedi=[c for c in root.iter('content') if c.get('startOffset') is not None and txt[int(c.get('startOffset')):int(c.get('startOffset'))+int(c.get('length'))]=='ivedi']
    assert ivedi and ivedi[0].get('bold')=='true','kalın sözcük korunur'
    assert page.locator('#partList li.missing').count()==0,page.locator('#partList li').all_text_contents()
    print('PASS yapıştır (Word): kaynak tanındı, sekme/numara/kalın korundu, görsel bildirildi, UDF standart')
    # 2) Google Dokümanlar: dış <b style=font-weight:normal> sarmalayıcısı kalın sayılmaz.
    paste_into(page,GDOCS_HTML,'x')
    assert 'Google' in page.locator('#pasteSource').text_content()
    body=page.locator('#paper .paragraph',has_text='ücret alacağı')
    assert body.locator('span',has_text='ayrıca').evaluate("s=>getComputedStyle(s).fontStyle")=='italic'
    assert body.locator('span',has_text='Davacı').evaluate("s=>getComputedStyle(s).fontWeight") in ('400','normal')
    assert page.locator('#paper [data-role="major-heading"]',has_text='AÇIKLAMALAR').count()==1,'eksik AÇIKLAMALAR eklendi'
    assert page.locator('#paper [data-role="result-heading"]').count()==1,'eksik SONUÇ VE İSTEM eklendi'
    print('PASS yapıştır (Google Dokümanlar): biçim doğru okundu, eksik bölüm başlıkları eklendi')
    # 3) PDF'den kopya: kırık satırlar birleşir.
    paste_into(page,'',PDF_TEXT)
    assert 'PDF' in page.locator('#pasteSource').text_content()
    assert page.locator('#paper .paragraph',has_text='14.10.2026 tarihli duruşmasına, şehir dışında').count()==1
    assert page.locator('#paper .paragraph',has_text='Sayfa 1').count()==0
    print('PASS yapıştır (PDF): satırlar birleşti, sayfa numarası atıldı')
    # 4) Yapay zekâ / Markdown yanıtı.
    paste_into(page,'','## MAZERET DİLEKÇESİ\n**ŞANLIURFA 8. ASLİYE HUKUK MAHKEMESİ\'NE**\n**DAVALI:** İsa ÇİFTÇİ\n**KONU:** Mazeret.\n\nDuruşmaya katılamayacağız.\n\nYukarıda açıklanan nedenlerle mazeretimizin kabulüne karar verilmesini talep ederim.')
    assert 'Yapay zekâ' in page.locator('#pasteSource').text_content()
    assert '**' not in page.locator('#paper').inner_text()
    assert page.locator('#paper [data-role="court-heading"]').count()==1
    print('PASS yapıştır (yapay zekâ/Markdown): işaretler temizlendi, dilekçe kuruldu')
    # 5) Dağınık düz metin + eksikleri formdan tamamla.
    paste_into(page,'','Şanlıurfa 8. Asliye Hukuk Mahkemesine\n\nDuruşmaya şehir dışında olmamız nedeniyle katılamayacağız.\nYukarıda açıklanan nedenlerle mazeretimizin kabulüne karar verilmesini saygıyla talep ederiz.')
    assert page.locator('#completePanel').is_visible()
    assert page.locator('.cf[data-part="court"]').is_hidden(),'küçük harfli mahkeme satırı kendiliğinden düzeltildi'
    assert page.locator('.cf[data-part="caseNo"]').is_visible() and page.locator('.cf[data-part="topic"]').is_visible()
    page.locator('#pc_caseNo').fill('2026/1929');page.locator('#pc_caseNo').dispatch_event('change')
    page.locator('#pc_client').fill('İsa ÇİFTÇİ');page.locator('#pc_client').dispatch_event('change')
    page.locator('#pc_other').fill('Yapı ve Kredi Bankası Anonim Şirketi');page.locator('#pc_other').dispatch_event('change')
    page.locator('#pc_topic').fill('Duruşmaya ilişkin mazeretimizin sunulmasıdır.');page.locator('#pc_topic').dispatch_event('change')
    page.wait_for_timeout(200)
    out=tmp/'plain-paste.udf';save_download(page,'#download',out)
    _,lines=udf_lines(out)
    for e in ['ŞANLIURFA 8. ASLİYE HUKUK MAHKEMESİNE','DOSYA NO\t: 2026/1929 Esas','DAVALI\t: İsa ÇİFTÇİ','VEKİLİ\t: Av. Fatih NAKIŞLI (UETS: 16615-16448-62422)','DAVACI\t: Yapı ve Kredi Bankası Anonim Şirketi',
              'KONU\t: Duruşmaya ilişkin mazeretimizin sunulmasıdır.','AÇIKLAMALAR\t:','SONUÇ VE İSTEM\t: Yukarıda arz edilen sebeplerle;','1.\tMazeretimizin kabulüne,',f'Karar verilmesini saygıyla talep ederiz. {today}','Davalı İsa ÇİFTÇİ Vekili','Av. Fatih NAKIŞLI']:
        assert e in lines,(e,lines)
    assert page.locator('#partList li.added').count()>=6,page.locator('#partList li').all_text_contents()
    # Paragraf düzenleyici yapıştırılan dilekçede de çalışır.
    page.locator('#paper .paragraph[data-role="explanation-body"]').first.click()
    page.locator('.edit-sheet .chip',has_text='Kalın').click();page.keyboard.press('Escape')
    assert page.locator('#paper .paragraph.edited').count()==1
    rp=tmp/'paste-rapor.txt';save_download(page,'#reportDownload',rp);rt=rp.read_text(encoding='utf-8')
    assert 'yapıştırılan metin (Düz metin)' in rt and 'Dilekçeye eklenen parçalar' in rt,rt[:400]
    # Kutudaki metni elle düzenleyip yeniden oluşturma.
    page.locator('#pasteBox').fill('DAVALI: Ali VELİ\nKısa açıklama.\nKabulüne karar verilmesini talep ederim.')
    page.locator('#pasteBuild').click();page.wait_for_timeout(200)
    assert 'Düz metin' in page.locator('#pasteSource').text_content() and page.locator('#paper .paragraph',has_text='Ali VELİ').count()==1
    assert not errors,errors
    print('PASS yapıştır (dağınık metin): eksikler formdan tamamlandı, dilekçe baştan kuruldu, düzenleyici ve rapor çalışıyor')
    ctx.close()

def v17_no_jump(browser,petition):
    """1.7: changing a setting or a control in the edit sheet must not move
    the control the user is working with (the reported 'jumps to the top')."""
    top=lambda page,sel: page.evaluate(f'Math.round(document.querySelector({sel!r}).getBoundingClientRect().top)')
    for w,h in [(390,844),(1440,900)]:
        ctx=browser.new_context(viewport={'width':w,'height':h});page=ctx.new_page()
        load_page(page,'workspace.html');fetch_data(page,petition)
        # Back to the legal profile: its long change list appears above the select.
        page.locator('#typographyProfile').select_option('faithful');page.wait_for_timeout(100)
        for sel,act in [('#typographyProfile',lambda: page.locator('#typographyProfile').select_option('legal')),
                        ('#splitRequests',lambda: page.locator('#splitRequests').click()),
                        ('#listMode',lambda: page.locator('#listMode').select_option('auto')),
                        ('#lineSpacing',lambda: page.locator('#lineSpacing').select_option('1'))]:
            page.locator(sel).evaluate("el=>el.scrollIntoView({block:'center'})");page.wait_for_timeout(80)
            before=top(page,sel);act();page.wait_for_timeout(200);after=top(page,sel)
            assert abs(after-before)<=4,(w,sel,before,after)
        # Edit sheet: its own scroll position and the focused control survive a change.
        page.locator('#paper .paragraph[data-role="explanation-body"]').first.click();page.wait_for_timeout(350)
        body='.edit-sheet .sheet-body'
        page.locator(body).evaluate('el=>el.scrollTop=el.scrollHeight');page.wait_for_timeout(80)
        before=page.locator(body).evaluate('el=>el.scrollTop')
        page.locator('.edit-sheet .chip',has_text='Önüne boş satır').click();page.wait_for_timeout(200)
        after=page.locator(body).evaluate('el=>el.scrollTop')
        assert before>50 and abs(after-before)<=2,(w,before,after)
        assert page.evaluate("document.activeElement.textContent")=='Önüne boş satır'
        ctx.close()
    print('PASS kaydırma sıçraması yok: ayar değişince kontrol yerinde kalır, düzenleme paneli kaydırması ve odak korunur')

SELECT_JS="""([sel,needle])=>{
  const p=[...document.querySelectorAll(sel)].find(x=>x.textContent.includes(needle));if(!p)return false;
  const w=document.createTreeWalker(p,NodeFilter.SHOW_TEXT);const nodes=[];let t,all='';
  while((t=w.nextNode())){nodes.push([t,all.length]);all+=t.data;}
  const i=all.indexOf(needle);const at=x=>{for(const [n,o] of nodes)if(x>=o&&x<=o+n.data.length)return [n,x-o];};
  const a=at(i),z=at(i+needle.length);const r=document.createRange();r.setStart(a[0],a[1]);r.setEnd(z[0],z[1]);
  const s=getSelection();s.removeAllRanges();s.addRange(r);return true;}"""

def udf_runs(path):
    """[(text, attributes)] for every content run of the UDF."""
    root=data_xml(path);txt=udf_text(root);out=[]
    for c in root.xpath('./elements/paragraph//content'):
        if c.get('startOffset') is None:continue
        a=int(c.get('startOffset'));out.append((txt[a:a+int(c.get('length'))].rstrip('\n'),dict(c.attrib)))
    return out

def v18_marks(browser,tmp,petition):
    """1.8: word/sentence formatting from a preview selection (format bar,
    Ctrl+B/I/U, sentence, everywhere), undo/redo, sentence buttons in the
    panel, new paragraph + move, phone docked bar, Word caps and spacing."""
    ctx=browser.new_context(accept_downloads=True,viewport={'width':1440,'height':900})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');fetch_data(page,petition)
    styled=lambda prop,val:page.evaluate(f"[...document.querySelectorAll('#paper span')].filter(s=>s.style.{prop}==='{val}').map(s=>s.textContent)")
    assert page.evaluate(SELECT_JS,['#paper .paragraph','verilen karara']);page.wait_for_timeout(250)
    assert page.locator('.mark-pop').is_visible(),'seçim çubuğu açılır'
    page.locator('.mark-pop .mp-bold').click();page.wait_for_timeout(150)
    assert 'verilen karara' in styled('fontWeight','700'),styled('fontWeight','700')
    assert page.evaluate('getSelection().toString()')=='verilen karara','seçim korunur'
    page.keyboard.press('Control+i');page.wait_for_timeout(150)
    assert 'verilen karara' in styled('fontStyle','italic')
    page.keyboard.press('Control+z');page.wait_for_timeout(150)
    assert 'verilen karara' not in styled('fontStyle','italic'),'geri al'
    page.keyboard.press('Control+y');page.wait_for_timeout(150)
    assert 'verilen karara' in styled('fontStyle','italic'),'yinele'
    # Cümleye genişlet + altı çizili.
    assert page.evaluate(SELECT_JS,['#paper .paragraph','şirketin merkezi']);page.wait_for_timeout(250)
    page.locator('.mark-pop .mp-sentence').click();page.wait_for_timeout(150)
    assert page.evaluate('getSelection().toString()')=='Müvekkil şirketin merkezi Gaziantep’tir.',page.evaluate('getSelection().toString()')
    page.locator('.mark-pop .mp-underline').click();page.wait_for_timeout(150)
    # Tüm belgede: her "ihtiyati haciz" (büyük/küçük harf duyarsız) kalın.
    assert page.evaluate(SELECT_JS,['#paper .paragraph[data-role="result-item"]','İhtiyati haciz']);page.wait_for_timeout(250)
    page.locator('.mark-pop .mp-scope').click();page.locator('.mark-pop .mp-bold').click();page.wait_for_timeout(200)
    trl=lambda t:t.replace('İ','i').replace('I','ı').lower()
    hits=[t for t in styled('fontWeight','700') if trl(t).startswith('ihtiyati haciz')]
    assert len(hits)>=2,hits
    page.evaluate('getSelection().removeAllRanges()');page.keyboard.press('Escape');page.wait_for_timeout(150)
    # Panel: cümle düğmesi → üstü çizili; biçim listesinden kaldır.
    page.locator('#paper .paragraph',has_text='Müvekkil şirketin').click();page.wait_for_timeout(300)
    assert page.locator('.edit-sheet .mark-list li').count()==1,'önizlemeden yapılan biçim panelde listelenir'
    page.locator('.edit-sheet .sent-chip').first.click();page.locator('.edit-sheet .mark-btn.mark-strike').click();page.wait_for_timeout(200)
    assert 'Müvekkil şirketin merkezi Gaziantep’tir.' in page.evaluate("[...document.querySelectorAll('#paper span')].filter(s=>s.style.textDecoration.includes('line-through')).map(s=>s.textContent)")
    # Aynı cümleye ikinci biçim, mevcut kaydıyla birleşir.
    assert page.locator('.edit-sheet .mark-list li').count()==1 and 'altı çizili + üstü çizili' in page.locator('.edit-sheet .mark-list li').text_content()
    page.locator('.edit-sheet .mark-list li').first.locator('button').click();page.wait_for_timeout(200)
    assert page.locator('.edit-sheet .mark-list li').count()==0
    assert page.locator('.edit-sheet .sent-chip.on').count()==1,'cümle seçili kalır'
    page.locator('.edit-sheet .mark-btn.mark-underline').click();page.wait_for_timeout(200)
    page.keyboard.press('Escape')
    # Yeni talep ekle (2. talebin altına) ve bir yukarı taşı.
    page.locator('#paper .paragraph[data-role="result-item"]',has_text='kaldırılmasına').click();page.wait_for_timeout(300)
    page.locator('.edit-sheet .chip',has_text='Altına yeni paragraf').click()
    page.locator('.edit-sheet .add-form textarea').fill('Teminatın iadesine,')
    page.locator('.edit-sheet .add-form button',has_text='Ekle').click();page.wait_for_timeout(300)
    assert page.locator('.edit-sheet .sheet-head strong').text_content()=='Eklenen paragraf'
    page.locator('.edit-sheet .chip',has_text='Yukarı taşı').click();page.wait_for_timeout(300)
    page.keyboard.press('Escape')
    udf=tmp/'v18.udf';save_download(page,'#download',udf)
    _,lines=udf_lines(udf)
    items=[l for l in lines if re.match(r'^\d\.\t',l)]
    assert items[:4]==['1.\tYetki itirazının kabulüne,','2.\tTeminatın iadesine,','3.\tİhtiyati haciz kararının kaldırılmasına,','4.\tYargılama giderlerinin karşı tarafa yükletilmesine,'],items
    runs=udf_runs(udf)
    assert any(t=='verilen karara' and a.get('bold')=='true' and a.get('italic')=='true' for t,a in runs),[r for r in runs if 'karara' in r[0]]
    assert any(t=='Müvekkil şirketin merkezi Gaziantep’tir.' and a.get('underline')=='true' for t,a in runs),[r for r in runs if 'Müvekkil' in r[0] or 'Gaziantep' in r[0]]
    assert sum(1 for t,a in runs if trl(t).startswith('ihtiyati haciz') and a.get('bold')=='true')>=2
    with page.expect_download() as info: page.locator('#reportDownload').click()
    rep=Path(info.value.path()).read_text(encoding='utf-8')
    assert '“verilen karara” kalın + italik' in rep and 'her “İhtiyati haciz” kalın' in rep and 'altına 1 yeni paragraf' in rep,rep[-1500:]
    assert not errors,errors
    ctx.close()
    # Telefon: seçim çubuğu altta sabit, ekranın içinde.
    ctx=browser.new_context(viewport={'width':390,'height':844},has_touch=True,is_mobile=True);page=ctx.new_page()
    load_page(page,'workspace.html');fetch_data(page,petition)
    page.locator('#paper .paragraph',has_text='verilen karara').scroll_into_view_if_needed()
    assert page.evaluate(SELECT_JS,['#paper .paragraph','verilen karara']);page.wait_for_timeout(300)
    box=page.locator('.mark-pop').bounding_box();bar=page.locator('#mobileBar').bounding_box()
    assert 'docked' in page.locator('.mark-pop').get_attribute('class') and box['y']+box['height']<=bar['y']+1 and box['x']>=0 and box['x']+box['width']<=390,(box,bar)
    page.locator('.mark-pop .mp-italic').click();page.wait_for_timeout(200)
    assert 'verilen karara' in styled('fontStyle','italic')
    ctx.close()
    # Word büyük harf / küçük büyük harf / karakter aralığı: metin büyük harfe, uyarı indirmeyi engellemez.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document()
    for t in ['ŞANLIURFA 1. ASLİYE HUKUK MAHKEMESİNE','DAVACI: Ali VELİ','KONU: Deneme.','AÇIKLAMALAR']:d.add_paragraph(t)
    p=d.add_paragraph();r=p.add_run('ilgili karar');r.font.all_caps=True;p.add_run(' metindedir; ');r=p.add_run('şerh');r.font.small_caps=True
    sp=OxmlElement('w:spacing');sp.set(qn('w:val'),'20');r2=p.add_run(' aralıklı.');r2._r.get_or_add_rPr().append(sp)
    capsdoc=tmp/'caps.docx';d.save(capsdoc)
    ctx=browser.new_context(accept_downloads=True);page=ctx.new_page();load_page(page,'workspace.html')
    page.locator('#fileInput').set_input_files(str(capsdoc));page.wait_for_timeout(1200)
    body=page.locator('#paper').text_content()
    assert 'İLGİLİ KARAR metindedir; ŞERH aralıklı.' in body,body
    assert page.locator('#reviewLabel').is_hidden() and page.locator('#download').is_enabled(),page.locator('#warningList li').all_text_contents()
    assert all(t.startswith('Bilgi:') for t in page.locator('#warningList li').all_text_contents() if 'aralığı' in t or 'small caps' in t)
    ctx.close()
    print('PASS kelime/cümle biçimi: seçim çubuğu, Ctrl+B/I/U, cümle, tüm belgede, geri al/yinele, panel cümleleri, yeni talep + taşı, telefon, UDF ve rapor; Word büyük harf efektleri')

V19_TEXT=("ŞANLIURFA 8. ASLİYE HUKUK MAHKEMESİNE\nDOSYA NO: 2026/1929 Esas\nDAVALI: İsa ÇİFTÇİ\nDAVACI: Yapı ve Kredi Bankası A.Ş.\n"
  "KONU: Mazeret.\nAÇIKLAMALAR\nMüvekkil İsa ÇİFTÇİ 20/10/2026 tarihli duruşmaya mazereti nedeniyle katılamayacaktır ,mazeret belgesi ektedir.\n"
  "SONUÇ VE İSTEM\n1. Mazeretimizin kabulüne,\nKarar verilmesini saygıyla talep ederiz.")

def v19_library(browser,tmp):
    """1.9: format suggestions with "Düzelt", find/replace (Ctrl+F, highlight,
    bold all, rules list), case cards and own templates across pages, export."""
    ctx=browser.new_context(accept_downloads=True,viewport={'width':1440,'height':900})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');paste_into(page,'',V19_TEXT);page.wait_for_timeout(300)
    checks=page.locator('#checkList li').all_text_contents()
    assert any(t.startswith('Öneri:') and 'GG/AA/YYYY' in t for t in checks),checks
    assert 'öneri' in page.locator('#checkStats').text_content()
    page.locator('#checkList li',has_text='GG/AA/YYYY').locator('button',has_text='Düzelt').click();page.wait_for_timeout(250)
    page.locator('#checkList li',has_text='önce boşluk').locator('button',has_text='Düzelt').click();page.wait_for_timeout(250)
    page.locator('#checkList li',has_text='virgülden sonra').locator('button',has_text='Düzelt').click();page.wait_for_timeout(250)
    paper=page.locator('#paper').text_content()
    assert '20.10.2026 tarihli' in paper and 'katılamayacaktır, mazeret belgesi' in paper,paper
    assert not any('GG/AA/YYYY' in t for t in page.locator('#checkList li').all_text_contents())
    # Bul / değiştir.
    page.locator('#paper').click(position={'x':5,'y':5});page.keyboard.press('Escape')
    page.keyboard.press('Control+f');page.wait_for_timeout(150)
    assert page.locator('.find-box').is_visible()
    page.locator('.find-in').fill('mazeret');page.wait_for_timeout(150)
    assert page.locator('.find-count').text_content().startswith('4 eşleşme'),page.locator('.find-count').text_content()
    assert page.evaluate("CSS.highlights?CSS.highlights.has('wtu-find'):true")
    page.locator('.find-box .find-bold').click();page.wait_for_timeout(200)
    bold=page.evaluate("[...document.querySelectorAll('#paper span')].filter(s=>s.style.fontWeight==='700').map(s=>s.textContent)")
    assert sum(1 for t in bold if t.lower()=='mazeret')>=3,bold
    page.locator('.find-in').fill('İsa ÇİFTÇİ');page.locator('.find-to').fill('İsa ÇİFTÇİ (müvekkil)');page.wait_for_timeout(100)
    page.locator('.find-box .find-replace').click();page.wait_for_timeout(250)
    assert page.locator('#paper').text_content().count('İsa ÇİFTÇİ (müvekkil)')>=2
    rules=page.locator('.rule-list li').all_text_contents()
    assert len(rules)==5 and any('İsa ÇİFTÇİ (müvekkil)' in r for r in rules),rules
    page.locator('.rule-list li',has_text='(müvekkil)').locator('button').click();page.wait_for_timeout(200)
    assert '(müvekkil)' not in page.locator('#paper').text_content(),'kural kaldırılınca metin geri döner'
    page.keyboard.press('Control+z');page.wait_for_timeout(200)
    assert '(müvekkil)' in page.locator('#paper').text_content(),'geri al kuralı geri getirir'
    page.locator('.find-box .find-close').click()
    # Dosya kartı ve şablon kaydet.
    page.locator('#libraryPanel summary').click()
    page.locator('#saveCaseFromDoc').click();page.wait_for_timeout(150)
    assert 'Dosya kartı kaydedildi: 2026/1929 Esas İsa ÇİFTÇİ' in page.locator('#status').text_content(),page.locator('#status').text_content()
    page.locator('#docTplName').fill('Mazeret – deneme');page.locator('#saveTplFromDoc').click();page.wait_for_timeout(150)
    assert 'şablon olarak kaydedildi' in page.locator('#status').text_content()
    assert 'Bu tarayıcıda 1 dosya kartı ve 1 kendi şablonunuz var.'==page.locator('#libraryStats').text_content()
    with page.expect_download() as info: page.locator('#exportSettings').click()
    exported=json.loads(Path(info.value.path()).read_text(encoding='utf-8'))
    assert exported['library']['cases'][0]['client'].startswith('İsa ÇİFTÇİ') and exported['library']['templates'][0]['title']=='Mazeret – deneme',exported['library']
    tpl=exported['library']['templates'][0]
    assert '{MÜVEKKİL}' in tpl['body'] and 'İsa ÇİFTÇİ' not in tpl['body'],tpl
    store=dump_storage(page);ctx.close()
    # Kayıtlı dosya kartı, yapıştırılan metindeki eksik künyeyi doldurur.
    ctx=browser.new_context(accept_downloads=True);page=ctx.new_page();load_page(page,'workspace.html',storage=store)
    paste_into(page,'','Duruşmaya katılamayacağız.\nYukarıda açıklanan nedenlerle mazeretimizin kabulüne karar verilmesini saygıyla talep ederiz.');page.wait_for_timeout(300)
    opts=page.locator('#pasteCasePicker select option').all_text_contents()
    assert len(opts)==2 and '2026/1929 Esas' in opts[1],opts
    page.locator('#pasteCasePicker select').select_option(index=1);page.wait_for_timeout(300)
    assert page.locator('#pc_caseNo').input_value()=='2026/1929 Esas' and page.locator('#pc_client').input_value().startswith('İsa ÇİFTÇİ')
    paper=page.locator('#paper').text_content()
    assert 'DOSYA NO: 2026/1929 Esas' in paper and 'DAVALI: İsa ÇİFTÇİ' in paper,paper[:300]
    ctx.close()
    # Şablonlar sayfası: kendi şablonu, alanlar, kayıt, silme.
    ctx=browser.new_context(accept_downloads=True);page=ctx.new_page();load_page(page,'templates.html',storage=store);page.wait_for_timeout(300)
    assert page.locator('#tpl optgroup').nth(1).get_attribute('label')=='Kendi şablonlarım (1)'
    page.locator('#tpl').select_option(label='Mazeret – deneme');page.wait_for_timeout(300)
    assert not page.locator('#tplDelete').is_hidden()
    page.locator('#casePicker select').select_option(index=1);page.wait_for_timeout(300)
    assert page.locator('#client').input_value().startswith('İsa ÇİFTÇİ')
    assert 'Müvekkil …' in page.locator('#body').input_value() or 'Müvekkil İsa' in page.locator('#body').input_value(),page.locator('#body').input_value()
    page.locator('#regen').click();page.wait_for_timeout(300)
    assert page.locator('#body').input_value().startswith('Müvekkil İsa ÇİFTÇİ'),page.locator('#body').input_value()
    page.locator('#body').fill('Davacıya {TUTAR} TL ödenmiştir.');page.locator('#tplName').fill('Ödeme beyanı');page.locator('#tplSave').click();page.wait_for_timeout(300)
    assert page.locator('#tpl').input_value().startswith('u:') and page.locator('#tplFields input').count()==1
    page.locator('#tplFields input').fill('5.000');page.wait_for_timeout(300)
    assert page.locator('#body').input_value()=='Davacıya 5.000 TL ödenmiştir.',page.locator('#body').input_value()
    assert 'Davacıya 5.000 TL ödenmiştir.' in page.locator('#paper').text_content()
    page.locator('#tplDelete').click();assert 'tekrar' in page.locator('#tplDelete').text_content()
    page.locator('#tplDelete').click();page.wait_for_timeout(300)
    assert page.locator('#tpl optgroup').nth(1).get_attribute('label')=='Kendi şablonlarım (1)','yalnız seçili şablon silinir'
    assert not errors,errors
    ctx.close()
    print('PASS kütüphane ve öneriler: Düzelt düğmeleri, Ctrl+F bul/değiştir ve vurgu, hepsi kalın, kural kaldır/geri al, dosya kartı ve kendi şablonu, sayfalar arası kullanım, dışa aktarma')

def v110_tools(browser,tmp,petition):
    """1.10: approximate pages and lonely signature, show changes, PDF print,
    file name pattern, experimental highlight, deadline page, new templates."""
    ctx=browser.new_context(accept_downloads=True,viewport={'width':1440,'height':900})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');fetch_data(page,petition)
    assert re.search(r'≈ \d+ sayfa',page.locator('#previewCount').text_content()),page.locator('#previewCount').text_content()
    # Sayfa hesabı: sentetik paragraf yükseklikleriyle; imza tek başına yeni sayfaya düşer.
    res=page.evaluate("""()=>{const paper=document.createElement('article');paper.className='paper';document.body.append(paper);
      const add=(h,role,text)=>{const p=document.createElement('div');p.className='paragraph';p.dataset.role=role;p.textContent=text;p.style.height=h+'px';paper.append(p);};
      add(994,'explanation-body','gövde');add(16,'result-closing','Karar verilmesini talep ederim.');add(30,'signature','Av. Ali VELİ');
      const r=WordToUDFPreview.paginate(paper,{margins:{top:42.52,bottom:42.52}});const g=paper.querySelectorAll('.page-guide').length;
      paper.lastChild.previousSibling.style.height='5px';paper.firstChild.style.height='900px';const r2=WordToUDFPreview.paginate(paper,{margins:{top:42.52,bottom:42.52}});
      paper.remove();return [r.pages,r.signatureAlone,g,r2.pages,r2.signatureAlone];}""")
    assert res==[2,True,1,1,False],res
    # Değişiklikleri göster.
    page.locator('#changesToggle').click();page.wait_for_timeout(250)
    marks=page.evaluate("[...document.querySelectorAll('#paper mark.chg')].map(m=>m.textContent)")
    assert marks and page.locator('#chgLegend').is_visible() and page.locator('#paper .has-chg').count()>=1,marks
    assert 'Dosya' not in marks,'kaynakta aynen olan sözcük işaretlenmez'
    page.locator('#changesToggle').click();page.wait_for_timeout(250)
    assert page.locator('#paper mark.chg').count()==0 and page.locator('#chgLegend').is_hidden()
    # PDF: yalnız dilekçe, @page kenar boşlukları, sonra her şey eski hâline.
    page.evaluate("""()=>{window.print=()=>{const s=document.adoptedStyleSheets.at(-1);window.__p={text:document.getElementById('docPrint').innerText,cls:document.body.className,title:document.title,
      page:[...s.cssRules].map(r=>r.cssText).join(' ')};window.dispatchEvent(new Event('afterprint'));};}""")
    page.locator('#printPdf').click();page.wait_for_timeout(200)
    pr=page.evaluate('window.__p')
    assert 'İSTANBUL 22. ASLİYE TİCARET MAHKEMESİNE' in pr['text'] and pr['cls']=='print-doc' and 'size: a4' in pr['page'].lower() and 'landscape' not in pr['page'].lower() and '42.5' in pr['page'] and pr['title']==petition.stem,pr
    assert 'Önizleme' not in pr['text'] and 'Çoklu seçim' not in pr['text'],'yalnız dilekçe yazdırılır'
    after=page.evaluate("[document.body.className,document.adoptedStyleSheets.length,document.getElementById('docPrint').children.length,document.title]")
    assert after==['',0,0,'Word → UDF Dönüştürücü'],after
    # Dosya adı kalıbı indirilen dosyaya uygulanır.
    page.locator('.file-name summary').click();page.locator('#fileNamePattern').fill('{ESAS}_{TÜR}')
    assert page.locator('#fileNamePreview').text_content()=='→ 2026-713 E. 2026-708 K._İHTİYATİ HACİZ KARARINA İTİRAZ VE İVEDİ İNCELEME TALEPLİDİR.udf',page.locator('#fileNamePreview').text_content()
    name=save_download(page,'#download',tmp/'v110.udf')
    assert name=='2026-713 E. 2026-708 K._İHTİYATİ HACİZ KARARINA İTİRAZ VE İVEDİ İNCELEME TALEPLİDİR.udf',name
    # Deneysel vurgu: kapalıyken düğme yok; açınca seçime uygulanır, UDF'de background.
    assert page.evaluate(SELECT_JS,['#paper .paragraph','verilen karara']);page.wait_for_timeout(250)
    assert page.locator('.mark-pop .mp-highlight').is_hidden()
    page.evaluate('getSelection().removeAllRanges()');page.locator('#experimental').check();page.wait_for_timeout(300)
    assert page.locator('#experimentBox').is_visible()
    assert page.evaluate(SELECT_JS,['#paper .paragraph','verilen karara']);page.wait_for_timeout(250)
    page.locator('.mark-pop .mp-highlight').click();page.wait_for_timeout(200)
    hu=tmp/'v110-hl.udf';save_download(page,'#download',hu)
    assert any(t=='verilen karara' and a.get('background') for t,a in udf_runs(hu)),[r for r in udf_runs(hu) if 'karara' in r[0]]
    sample=tmp/'deneme.udf';save_download(page,'#experimentSample',sample)
    assert sum(1 for t,a in udf_runs(sample) if a.get('background'))==3
    assert not errors,errors
    ctx.close()
    # Süre hesapla sayfası ve takvim dosyası.
    ctx=browser.new_context(accept_downloads=True);page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'sure.html')
    page.locator('#dStart').fill('2026-07-10');page.locator('#dKind').select_option('istinaf');page.wait_for_timeout(150)
    assert page.locator('#dDate').text_content()=='7 Eylül 2026 Pazartesi' and page.locator('#dRecess').is_checked()
    page.locator('#dKind').select_option('odeme');page.wait_for_timeout(150)
    assert not page.locator('#dRecess').is_checked() and 'İcra' in page.locator('#dKindNote').text_content()
    assert page.locator('#dDate').text_content()=='17 Temmuz 2026 Cuma'
    page.locator('#dKind').select_option('custom');page.locator('#dAmount').fill('1');page.locator('#dUnit').select_option('ay');page.locator('#dStart').fill('2026-01-31');page.wait_for_timeout(150)
    assert page.locator('#dDate').text_content()=='2 Mart 2026 Pazartesi',page.locator('#dDate').text_content()
    page.locator('#dLabel').fill('2026/1929 E. – deneme')
    with page.expect_download() as info: page.locator('#dIcs').click()
    ics=Path(info.value.path()).read_text(encoding='utf-8')
    assert 'DTSTART;VALUE=DATE:20260302' in ics and '2026/1929 E. – deneme' in ics,ics
    assert not errors,errors
    ctx.close()
    # Yeni şablonlar: cevaba cevap / ikinci cevap tarafa göre.
    ctx=browser.new_context();page=ctx.new_page();load_page(page,'templates.html');page.wait_for_timeout(300)
    page.locator('#tpl').select_option('replik');page.wait_for_timeout(250)
    page.locator('#role').select_option('DAVACI');page.wait_for_timeout(300)
    assert 'cevaba cevap dilekçemizin' in page.locator('#topic').input_value()
    page.locator('#role').select_option('DAVALI');page.wait_for_timeout(300)
    assert 'ikinci cevap' in page.locator('#topic').input_value() and 'Davanın reddine' not in page.locator('#requests').input_value()
    assert 'davanın reddine' in page.locator('#requests').input_value().lower()
    ctx.close()
    print('PASS 1.10 araçları: yaklaşık sayfa ve yalnız kalan imza, değişiklikleri göster, PDF (yalnız dilekçe, @page), dosya adı kalıbı, deneysel vurgu (UDF background + deneme dosyası), süre hesabı ve takvim, cevaba cevap / ikinci cevap')

def v1101_audit(browser,tmp,petition):
    """1.10.1 denetim düzeltmeleri, tarayıcıda: klavyeyle paragraf, Türkçe Q
    Ctrl+I, üç tık, panel taslağı, tüm belgede tam sözcük, taşımaların
    birleşmesi, numaranın metne yapışmaması, yabancı JSON'un reddi, kendi
    şablonunun alanları, taraf varsayılanı, süre sayfası, Word stil uyarıları,
    DOCX öğe sırası, yerel verileri silme."""
    import json as _j
    ctx=browser.new_context(accept_downloads=True,viewport={'width':1440,'height':900})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');fetch_data(page,petition)
    # Klavye: Tab ile paragrafa gelinir, Enter paneli açar, Escape kapatır ve odak paragrafa döner.
    para=page.locator('#paper .paragraph[data-key]',has_text='verilen karara')
    assert para.get_attribute('tabindex')=='0'
    para.focus();page.keyboard.press('Enter');page.wait_for_timeout(250)
    assert page.locator('.edit-sheet').is_visible(),'Enter paneli açar'
    page.keyboard.press('Escape');page.wait_for_timeout(200)
    assert page.evaluate("document.activeElement?.dataset?.key")==para.get_attribute('data-key'),'odak paragrafa döner'
    # Türkçe Q klavye: Ctrl + ı italik yapar.
    assert page.evaluate(SELECT_JS,['#paper .paragraph','verilen karara']);page.wait_for_timeout(250)
    page.evaluate("document.dispatchEvent(new KeyboardEvent('keydown',{key:'ı',ctrlKey:true,bubbles:true,cancelable:true}))");page.wait_for_timeout(200)
    assert 'verilen karara' in page.evaluate("[...document.querySelectorAll('#paper span')].filter(s=>s.style.fontStyle==='italic').map(s=>s.textContent)")
    page.evaluate('getSelection().removeAllRanges()');page.keyboard.press('Escape')
    # Üç tık: paragrafın tamamı seçilir ve biçim çubuğu açılır (seçim sonraki paragrafın başında bitse de).
    page.locator('#paper .paragraph',has_text='Müvekkil şirketin merkezi').click(click_count=3);page.wait_for_timeout(350)
    assert page.locator('.mark-pop').is_visible(),'üç tıkla biçim çubuğu'
    page.locator('.mark-pop .mp-underline').click();page.wait_for_timeout(200)
    assert 'Müvekkil şirketin merkezi Gaziantep’tir.' in page.evaluate("[...document.querySelectorAll('#paper span')].filter(s=>s.style.textDecoration.includes('underline')).map(s=>s.textContent)")
    page.keyboard.press('Control+z');page.wait_for_timeout(150)
    page.evaluate('getSelection().removeAllRanges()');page.keyboard.press('Escape');page.wait_for_timeout(150)
    # Panelde yazılan ama uygulanmayan metin, başka bir düğmeye basınca kaybolmaz.
    page.locator('#paper .paragraph',has_text='Müvekkil şirketin merkezi').click();page.wait_for_timeout(300)
    page.locator('.edit-sheet .sheet-text').fill('Taslak metin, henüz uygulanmadı.')
    page.locator('.edit-sheet .chip.style-b').click();page.wait_for_timeout(250)
    assert page.locator('.edit-sheet .sheet-text').input_value()=='Taslak metin, henüz uygulanmadı.'
    page.keyboard.press('Escape')
    # Word → DOCX: rPr öğeleri şema sırasında (b, i … sz … u).
    dx=tmp/'v1101.docx';save_download(page,'#downloadDocx',dx)
    with zipfile.ZipFile(dx) as z:xml=z.read('word/document.xml').decode('utf-8')
    for rpr in re.findall(r'<w:rPr>(.*?)</w:rPr>',xml):
        tags=re.findall(r'<w:(\w+)',rpr);order=['rFonts','b','bCs','i','iCs','strike','color','sz','szCs','highlight','u','shd']
        idx=[order.index(t) for t in tags if t in order];assert idx==sorted(idx),tags
    for ppr in re.findall(r'<w:pPr>(.*?)</w:pPr>',xml):
        tags=[t for t in re.findall(r'<w:(\w+)',ppr) if t in ('pageBreakBefore','numPr','tabs','spacing','ind','jc')];order=['pageBreakBefore','numPr','tabs','spacing','ind','jc']
        assert [order.index(t) for t in tags]==sorted(order.index(t) for t in tags),tags
    # WordToUDF'ye ait olmayan JSON ayarları sıfırlamaz.
    before=page.locator('#ownName').input_value()
    bad=tmp/'paket.json';bad.write_text('{"name":"paket","version":"1.0.0"}',encoding='utf-8')
    page.locator('#importSettings').set_input_files(str(bad));page.wait_for_timeout(300)
    assert 'WordToUDF ayar dosyası değil' in page.locator('#status').text_content() and page.locator('#ownName').input_value()==before
    assert not errors,errors
    ctx.close()
    # Yapıştırılan talepler: bir sözcüğü kalın yapıp 1. maddeyi silince numaralar 1., 2.; tüm belgede tam sözcük; taşımalar.
    ctx=browser.new_context(viewport={'width':1440,'height':900});page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html')
    html=('<p>ANKARA 1. ASLİYE HUKUK MAHKEMESİNE</p><p>DAVACI: Ali VELİ</p><p>KONU: Deneme.</p><p>AÇIKLAMALAR</p>'
          '<p>Paragraf A: Ali talimat verdi.</p><p>Paragraf B: Halil geldi.</p><p>Paragraf C: Ali gitti.</p><p>Paragraf D: son.</p>'
          '<p>SONUÇ VE İSTEM</p><p>1. Mazeretimizin kabulüne,</p><p>2. Duruşmanın ertelenmesine,</p><p>3. Giderlerin yükletilmesine,</p>')
    paste_into(page,html,'x');page.wait_for_timeout(300)
    assert page.evaluate(SELECT_JS,['#paper .paragraph','ertelenmesine']);page.wait_for_timeout(250)
    page.locator('.mark-pop .mp-bold').click();page.wait_for_timeout(200)
    page.evaluate('getSelection().removeAllRanges()');page.keyboard.press('Escape')
    page.locator('#paper .paragraph',has_text='Mazeretimizin kabulüne').click();page.wait_for_timeout(250)
    page.locator('.edit-sheet .chip',has_text='Paragrafı sil').click();page.wait_for_timeout(250);page.keyboard.press('Escape')
    items=[t.replace('\u00a0',' ') for t in page.locator('#paper .paragraph[data-role="result-item"]').all_text_contents()]
    assert len(items)==2 and items[0].startswith('1.') and 'Duruşmanın' in items[0] and items[1].startswith('2.'),items
    # Tüm belgede: "Ali" yalnız tam sözcük olarak (talimat, Halil değil).
    assert page.evaluate(SELECT_JS,['#paper .paragraph','Ali gitti']);page.wait_for_timeout(100)
    page.evaluate(SELECT_JS,['#paper .paragraph','Ali gitti'])
    page.evaluate("""()=>{const p=[...document.querySelectorAll('#paper .paragraph')].find(x=>x.textContent.includes('Ali gitti'));const w=document.createTreeWalker(p,NodeFilter.SHOW_TEXT);let t;while((t=w.nextNode())){const i=t.data.indexOf('Ali gitti');if(i>=0){const r=document.createRange();r.setStart(t,i);r.setEnd(t,i+3);const s=getSelection();s.removeAllRanges();s.addRange(r);return;}}}""")
    page.wait_for_timeout(250)
    page.locator('.mark-pop .mp-scope').click();page.locator('.mark-pop .mp-bold').click();page.wait_for_timeout(250)
    bold=page.evaluate("[...document.querySelectorAll('#paper span')].filter(s=>s.style.fontWeight==='700').map(s=>s.textContent)")
    assert 'Ali' in bold and not any(t in ('ali','Hal') for t in bold) and not any('talimat' in t for t in bold),bold
    page.evaluate('getSelection().removeAllRanges()');page.keyboard.press('Escape')
    # Taşıma: C yukarı, sonra B yukarı → A B C D.
    for name in ['Paragraf C','Paragraf B']:
        page.locator('#paper .paragraph',has_text=name).click();page.wait_for_timeout(250)
        page.locator('.edit-sheet .chip',has_text='Yukarı taşı').click();page.wait_for_timeout(250);page.keyboard.press('Escape');page.wait_for_timeout(150)
    order=[t[9] for t in page.locator('#paper .paragraph').all_text_contents() if t.startswith('Paragraf ')]
    assert order==['A','B','C','D'],order
    assert not errors,errors
    ctx.close()
    # Kendi şablonu: aynı adla yeniden kaydetmek alanları korur; düzenlenmiş metinde değerler alana döner.
    ctx=browser.new_context(accept_downloads=True);page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'templates.html');page.wait_for_timeout(300)
    page.locator('#tpl').select_option('feragat');page.wait_for_timeout(250)
    assert page.locator('#role').input_value()=='DAVACI','feragat davacı tarafından yapılır'
    page.locator('#tpl').select_option('beyan');page.wait_for_timeout(250)
    page.locator('#client').fill('Ali VELİ')
    page.locator('#body').fill('Müvekkil {MÜVEKKİL} {TARİH} tarihinde {ALACAK TUTARI} TL ödedi.');page.locator('#tplName').fill('Ödeme');page.locator('#tplSave').click();page.wait_for_timeout(300)
    page.locator('#tplFields input').fill('10.000');page.wait_for_timeout(300)
    assert 'Ali VELİ' in page.locator('#body').input_value() and '10.000' in page.locator('#body').input_value()
    page.locator('#tplName').fill('Ödeme');page.locator('#tplSave').click();page.wait_for_timeout(300)
    tpl=[t for t in _j.loads(dump_storage(page)['wordtoudf.userTemplates.v1']) if t['title']=='Ödeme'][0]
    assert tpl['body']=='Müvekkil {MÜVEKKİL} {TARİH} tarihinde {ALACAK TUTARI} TL ödedi.',tpl['body']
    page.locator('#tplFields input').fill('10.000');page.wait_for_timeout(300)
    page.locator('#body').fill(page.locator('#body').input_value()+' Ödeme dekontu ektedir.')
    page.locator('#tplName').fill('Ödeme');page.locator('#tplSave').click();page.wait_for_timeout(300)
    tpl=[t for t in _j.loads(dump_storage(page)['wordtoudf.userTemplates.v1']) if t['title']=='Ödeme'][0]
    assert tpl['body']=='Müvekkil {MÜVEKKİL} {TARİH} tarihinde {ALACAK TUTARI} TL ödedi. Ödeme dekontu ektedir.',tpl['body']
    assert not errors,errors
    store=dump_storage(page);ctx.close()
    # Yerel verileri silme: iki basış; ayarlar kalır.
    store['wordtoudf.cases.v1']=_j.dumps([{'id':'abcdef12','court':'Ankara 1. Asliye Hukuk Mahkemesi','caseNo':'2026/1','client':'Ali VELİ','updated':'2026-10-01T00:00:00.000Z'}])
    ctx=browser.new_context();page=ctx.new_page();load_page(page,'workspace.html',storage=store)
    page.locator('#libraryPanel summary').click()
    assert 'Bu tarayıcıda 1 dosya kartı ve 1 kendi şablonunuz var.'==page.locator('#libraryStats').text_content(),page.locator('#libraryStats').text_content()
    page.locator('#clearLibrary').click();assert 'tekrar' in page.locator('#clearLibrary').text_content()
    page.locator('#clearLibrary').click();page.wait_for_timeout(200)
    left=dump_storage(page)
    assert not any(k in left for k in ['wordtoudf.cases.v1','wordtoudf.userTemplates.v1','wordtoudf.templateDraft.v1']),list(left)
    assert page.locator('#libraryStats').text_content()=='Bu tarayıcıda 0 dosya kartı ve 0 kendi şablonunuz var.'
    ctx.close()
    # Süre sayfası: geçersiz miktar eski sonucu bırakmaz; birime göre sınır.
    ctx=browser.new_context();page=ctx.new_page();load_page(page,'sure.html')
    page.locator('#dKind').select_option('custom');page.locator('#dStart').fill('2026-07-10');page.wait_for_timeout(100)
    assert page.locator('#dSteps li').count()>0
    page.locator('#dAmount').fill('');page.wait_for_timeout(100)
    assert page.locator('#dDate').text_content()=='—' and page.locator('#dLeft').text_content()=='' and page.locator('#dSteps li').count()==0
    page.locator('#dAmount').fill('11');page.locator('#dUnit').select_option('yil');page.wait_for_timeout(100)
    assert '1 ile 10 yıl' in page.locator('#status').text_content()
    page.locator('#dAmount').fill('2');page.locator('#dUnit').select_option('hafta');page.locator('#dStart').fill('2025-08-16');page.wait_for_timeout(100)
    assert page.locator('#dDate').text_content()=='8 Eylül 2025 Pazartesi',page.locator('#dDate').text_content()
    ctx.close()
    # Word: kullanılmayan üst simge stili engelleyici uyarı vermez; gerçek üst simge verir.
    from docx.enum.style import WD_STYLE_TYPE
    d=Document()
    st=d.styles.add_style('Dipnot Başvurusu',WD_STYLE_TYPE.CHARACTER);st.font.superscript=True
    for t in ['ANKARA 1. ASLİYE HUKUK MAHKEMESİNE','DAVACI: Ali VELİ','KONU: Deneme.','AÇIKLAMALAR','Metin.']:d.add_paragraph(t)
    plain=tmp/'unused-style.docx';d.save(plain)
    p2=d.add_paragraph('Alan 10 m');r=p2.add_run('2');r.font.superscript=True
    sup=tmp/'superscript.docx';d.save(sup)
    ctx=browser.new_context();page=ctx.new_page();load_page(page,'workspace.html')
    page.locator('#fileInput').set_input_files(str(plain));page.wait_for_timeout(1200)
    warns=' '.join(page.locator('#warningList li').all_text_contents())
    assert 'simge' not in warns and 'karakter aralığı' not in warns and page.locator('#reviewLabel').is_hidden(),warns
    page.locator('#fileInput').set_input_files(str(sup));page.wait_for_timeout(1200)
    assert 'simge' in ' '.join(page.locator('#warningList li').all_text_contents())
    ctx.close()
    # Açılır pencere doğru şablon sayısını söyler.
    assert '14 hazır dilekçe şablonu' in (ROOT/'popup.html').read_text(encoding='utf-8')
    print('PASS 1.10.1 denetim: klavyeyle paragraf, Ctrl+ı, üç tık, panel taslağı, DOCX öğe sırası, yabancı JSON reddi, numara metne yapışmaz, tüm belgede tam sözcük, taşımalar birleşir, şablon alanları korunur, feragat davacı, yerel verileri silme, süre sayfası, kullanılmayan Word stili')

def v111_design(browser,tmp,petition):
    """1.11: phone preview fits the text column and opens full screen (zoom,
    editing above the bottom sheet, Esc), command palette, step tracker,
    check ring with fixed suggestions and their undo, template gallery,
    deadline e-Tebligat and calendar."""
    # --- Telefon: metin sütunu ekranı doldurur; tam ekran, yakınlaştırma, panelin üstünde düzenleme.
    ctx=browser.new_context(viewport={'width':390,'height':844},has_touch=True,is_mobile=True)
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');fetch_data(page,petition)
    scale=page.evaluate("()=>parseFloat(document.getElementById('zoomToggle').textContent.match(/%(\\d+)/)[1])/100")
    whole=page.evaluate("()=>{const v=document.getElementById('previewViewport'),s=getComputedStyle(v);return (v.clientWidth-parseFloat(s.paddingLeft)-parseFloat(s.paddingRight))/794}")
    assert scale>whole+0.04,(scale,whole)
    sw=page.evaluate("()=>[document.documentElement.scrollWidth,innerWidth,document.getElementById('previewViewport').scrollWidth,document.getElementById('previewViewport').clientWidth]")
    assert sw[0]<=sw[1]+1 and sw[2]<=sw[3]+1,sw
    page.locator('#mbPreview').click();page.wait_for_timeout(250)
    assert page.evaluate("getComputedStyle(document.getElementById('previewSection')).position")=='fixed' and page.locator('#previewFull').get_attribute('aria-pressed')=='true'
    assert page.locator('#tabbar, .tabbar').first.is_hidden() and page.locator('#mobileBar').is_hidden()
    page.locator('#zoomIn').click();page.locator('#zoomIn').click();page.wait_for_timeout(100)
    z=page.locator('#zoomToggle').text_content();assert re.fullmatch(r'%\d+',z),z
    page.locator('#zoomToggle').click();page.wait_for_timeout(100);assert 'Sığdır' in page.locator('#zoomToggle').text_content()
    page.locator('#paper .paragraph[data-role="result-item"]').last.click();page.wait_for_timeout(700)
    pos=page.evaluate("()=>{const p=document.querySelector('#paper .paragraph.selected').getBoundingClientRect(),s=document.querySelector('.edit-sheet');return [Math.round(p.top),Math.round(p.bottom),innerHeight-s.offsetHeight]}")
    assert pos[0]>=0 and pos[1]<=pos[2],('düzenlenen paragraf alttaki panelin üstünde görünür',pos)
    page.locator('.edit-sheet .chip',has_text='Kalın').click();page.wait_for_timeout(150)
    assert page.locator('#paper .paragraph.edited').count()==1
    page.keyboard.press('Escape');page.wait_for_timeout(100);assert page.locator('.edit-sheet').count()==0
    assert page.evaluate("getComputedStyle(document.getElementById('previewSection')).position")=='fixed','önce panel kapanır'
    page.keyboard.press('Escape');page.wait_for_timeout(100)
    assert page.evaluate("getComputedStyle(document.getElementById('previewSection')).position")!='fixed' and page.locator('#previewFull').get_attribute('aria-pressed')=='false'
    # Satır içi önizlemede de paragraf panelin üstüne gelir.
    page.locator('#paper .paragraph[data-role="result-item"]').first.click();page.wait_for_timeout(700)
    pos=page.evaluate("()=>{const p=document.querySelector('#paper .paragraph.selected').getBoundingClientRect(),s=document.querySelector('.edit-sheet');return [Math.round(p.top),Math.round(p.bottom),innerHeight-s.offsetHeight]}")
    assert pos[0]>=0 and pos[1]<=pos[2],pos
    assert not errors,errors
    ctx.close()
    # --- Masaüstü: komut paleti, adımlar, denetim halkası ve düzeltilenlerin geri alınması.
    ctx=browser.new_context(accept_downloads=True,viewport={'width':1440,'height':900})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html')
    assert page.evaluate("[...document.querySelectorAll('#steps li')].map(l=>l.className.trim())")[0]=='active'
    paste_into(page,'',V19_TEXT);page.wait_for_timeout(300)
    st=page.evaluate("[...document.querySelectorAll('#steps li')].map(l=>l.className.trim())")
    assert st[0]=='done' and 'active' in st[3],st
    assert page.locator('#checkRing .ring-seg').count()>=2 and page.locator('#ringBig').text_content().isdigit()
    n=page.locator('#checkRing .ring-seg').count()
    page.locator('#checkList li',has_text='GG/AA/YYYY').locator('button',has_text='Düzelt').click();page.wait_for_timeout(250)
    fixed=page.locator('#checkList li.check-fixed')
    assert fixed.count()==1 and '20.10.2026' in page.locator('#paper').text_content(),page.locator('#checkList li').all_text_contents()
    assert page.locator('#checkRing .seg-fixed').count()==1 and page.locator('#checkRing .ring-seg').count()==n
    fixed.locator('button',has_text='Geri al').click();page.wait_for_timeout(250)
    assert page.locator('#checkList li.check-fixed').count()==0 and '20/10/2026' in page.locator('#paper').text_content()
    page.keyboard.press('Control+k');page.wait_for_timeout(100)
    assert page.locator('.palette').count()==1 and page.evaluate("document.activeElement.id")=='palInput'
    page.keyboard.type('bul değiş');page.wait_for_timeout(100)
    assert 'Bul / değiştir' in page.locator('.pal-item[aria-selected="true"]').text_content()
    page.keyboard.press('Enter');page.wait_for_timeout(150)
    assert page.locator('.palette').count()==0 and page.locator('.find-box').is_visible()
    page.keyboard.press('Control+k');page.keyboard.type('mazeret');page.wait_for_timeout(100)
    assert page.locator('.pal-item').count()>=1 and 'Duruşma mazeret' in page.locator('.pal-item').first.text_content()
    page.keyboard.press('Escape');assert page.locator('.palette').count()==0
    # İndirme düğmesi: indirilince yeşil "İndirildi", adım 4 tamam.
    save_download(page,'#download',tmp/'v111.udf');page.wait_for_timeout(100)
    assert page.locator('#download').get_attribute('data-state')=='done' and 'İndirildi' in page.locator('#download').text_content()
    assert all('done' in c for c in page.evaluate("[...document.querySelectorAll('#steps li')].map(l=>l.className)"))
    assert not errors,errors
    ctx.close()
    # --- Şablon galerisi: kart seçimi #tpl ile aynı; arama; kategori.
    ctx=browser.new_context(viewport={'width':1440,'height':900});page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'templates.html');page.wait_for_timeout(300)
    assert page.locator('#tplGrid .tpl-card').count()==14
    page.locator('#tplSearch').fill('bilirkişi');page.wait_for_timeout(100)
    titles=page.locator('#tplGrid .tpl-card strong').all_text_contents();assert titles and all('ilirkişi' in t for t in titles),titles
    page.locator('#tplGrid .tpl-card',has_text='Bilirkişi raporuna itiraz').click();page.wait_for_timeout(300)
    assert page.locator('#tpl').input_value()=='bilirkisi' and page.locator('#tplTitle').text_content().startswith('Bilirkişi')
    assert 'ilirkişi raporu' in page.locator('#paper').text_content()
    assert not page.locator('#tplGallery').evaluate('d=>d.open'),'seçimden sonra galeri katlanır, form öne gelir'
    page.locator('#tplGallery summary').click();page.locator('#tplSearch').fill('');page.locator('#tplCats .chip',has_text='Duruşma').click();page.wait_for_timeout(100)
    assert page.locator('#tplGrid .tpl-card').count()==4
    page.locator('#tpl').select_option('islah');page.wait_for_timeout(250)
    page.locator('#tplCats .chip',has_text='Hepsi').click();page.wait_for_timeout(100)
    assert page.locator('#tplGrid .tpl-card[aria-pressed="true"]').get_attribute('data-tpl')=='islah','seçim kutusu galeriyle aynı'
    assert not errors,errors
    ctx.close()
    # --- Süre: e-Tebligat +5 gün ve takvim.
    ctx=browser.new_context(viewport={'width':1440,'height':900});page=ctx.new_page();load_page(page,'sure.html')
    page.locator('#dStart').fill('2026-10-15');page.locator('.d-preset',has_text='Cevap').first.click();page.wait_for_timeout(150)
    assert page.locator('#dDate').text_content()=='30 Ekim 2026 Cuma',page.locator('#dDate').text_content()  # 29 Ekim tatil
    page.locator('#dEtebligat').check();page.wait_for_timeout(150)
    assert page.locator('#dDate').text_content()=='3 Kasım 2026 Salı',page.locator('#dDate').text_content()
    assert 'Tebligat K. m. 7/a' in page.locator('#dSteps li').first.text_content()
    assert page.locator('#dMonths .cal-day.end').count()==1 and page.locator('#dMonths .cal-day.start').text_content()=='20'
    assert page.locator('#dMonths .cal-day.holiday').count()>=1  # 29 Ekim
    ctx.close()
    print('PASS 1.11 tasarım: telefonda metin sütunu ekranı doldurur, tam ekran + yakınlaştırma + panelin üstünde düzenleme, Esc sırası; adımlar, denetim halkası, düzeltileni geri al, komut paleti, indirme durumu, şablon galerisi, e-Tebligat ve takvim')

def v111_editing(browser,tmp,petition):
    """1.11 audit: a late web font does not rebuild the preview (the selection
    stays); the phone's format bar stays on one row and toasts move above it;
    a paragraph moved up in full screen stays visible under the tool row; on
    a desktop the full-screen preview starts right of the edit panel."""
    ctx=browser.new_context(viewport={'width':390,'height':844},has_touch=True,is_mobile=True)
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');fetch_data(page,petition)
    page.locator('#paper .paragraph',has_text='verilen karara').scroll_into_view_if_needed()
    assert page.evaluate(SELECT_JS,['#paper .paragraph','verilen karara']);page.wait_for_timeout(400)
    page.evaluate("document.querySelector('#paper .paragraph[data-key]').__keep=1")
    # Yazı tipi geç yüklendi: Tinos ölçüyü yeniler ama belgeyi baştan çizmez; arayüz yazı tipi hiç dokunmaz.
    for fam in ('Manrope','Tinos'):
        page.evaluate("f=>document.fonts.dispatchEvent(new FontFaceSetLoadEvent('loadingdone',{fontfaces:[new FontFace(f,'local(Arial)')]}))",fam);page.wait_for_timeout(250)
    assert page.evaluate("document.querySelector('#paper .paragraph[data-key]').__keep===1"),'önizleme yeniden çizilmedi'
    assert page.evaluate('getSelection().toString()')=='verilen karara' and page.locator('.mark-pop').is_visible(),'seçim ve biçim çubuğu kalır'
    bar=page.locator('.mark-pop').bounding_box();assert bar['height']<=60 and bar['x']>=0 and bar['x']+bar['width']<=390,bar
    assert page.evaluate("(p=>p.scrollWidth<=p.clientWidth+1)(document.querySelector('.mark-pop'))"),'bütün düğmeler görünür (Tümü)'
    page.locator('.mark-pop .mp-bold').click();page.wait_for_timeout(300)
    t=page.locator('.toast:not(.leaving)').bounding_box();bar=page.locator('.mark-pop').bounding_box()
    assert t and (t['y']+t['height']<=bar['y'] or t['y']>=bar['y']+bar['height']),('bildirim biçim çubuğunu örtmez',t,bar)
    # Tam ekranda yukarı taşınan paragraf araç satırının altına kaçmaz.
    page.keyboard.press('Escape');page.locator('#mbPreview').click();page.wait_for_timeout(300)
    page.locator('#paper .paragraph[data-role="result-item"]').last.click();page.wait_for_timeout(700)
    VIS="()=>{const p=document.querySelector('#paper .paragraph.selected').getBoundingClientRect(),v=document.getElementById('previewViewport').getBoundingClientRect(),s=document.querySelector('.edit-sheet');return p.top>=v.top-1&&p.bottom<=innerHeight-s.offsetHeight+1}"
    for _ in range(8):
        up=page.locator('.edit-sheet .chip',has_text='Yukarı taşı')
        if up.is_disabled():break
        up.click();page.wait_for_timeout(500)
        assert page.evaluate(VIS),'taşınan paragraf görünür kalır'
    assert not errors,errors
    ctx.close()
    ctx=browser.new_context(viewport={'width':1280,'height':800})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');fetch_data(page,petition)
    page.locator('#previewFull').click();page.wait_for_timeout(300)
    page.locator('#paper .paragraph[data-role="explanation-body"]').first.click();page.wait_for_timeout(600)
    cover=page.evaluate("""()=>{const s=document.querySelector('.edit-sheet').getBoundingClientRect(),p=document.querySelector('#paper .paragraph.selected').getBoundingClientRect();
      const tools=[...document.querySelectorAll('#previewTools button')].filter(b=>b.getClientRects().length).some(b=>{const q=b.getBoundingClientRect();return q.left<s.right&&q.right>s.left&&q.top<s.bottom&&q.bottom>s.top});
      return {para:Math.max(0,Math.min(p.right,s.right)-Math.max(p.left,s.left)),tools}}""")
    assert cover=={'para':0,'tools':False},('panel sayfayı ve araçları örtmez',cover)
    page.keyboard.press('Escape');page.wait_for_timeout(200)
    assert page.evaluate("document.getElementById('previewSection').style.paddingLeft")=='','panel kapanınca önizleme yerine döner'
    assert not errors,errors
    ctx.close()
    # UDF → Word: telefonda sayfa ekrana sığar (yan kaydırma yok), boş sayfa yazısı görünür, tam ekran açılır.
    ctx=browser.new_context(viewport={'width':390,'height':844},has_touch=True,is_mobile=True)
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'reverse.html')
    box=page.locator('#paper .empty-state strong').bounding_box();assert box['x']>=0 and box['x']+box['width']<=390,box
    page.locator('#fileInput').set_input_files(str(tmp/'v111.udf'));page.wait_for_timeout(800)
    assert re.search(r'Sığdır · %\d+',page.locator('#zoomToggle').text_content())
    sw=page.evaluate("()=>{const v=document.getElementById('previewViewport');return [document.documentElement.scrollWidth,v.scrollWidth,v.clientWidth]}")
    assert sw[0]<=391 and sw[1]<=sw[2]+1,sw
    page.locator('#previewFull').click();page.wait_for_timeout(200)
    assert page.evaluate("getComputedStyle(document.getElementById('previewSection')).position")=='fixed'
    page.keyboard.press('Escape');page.wait_for_timeout(100)
    assert not errors,errors
    ctx.close()
    print('PASS 1.11 düzenleme: geç yüklenen yazı tipi önizlemeyi yeniden kurmaz (seçim kalır), telefonda biçim çubuğu tek satır ve bildirim üstünde, tam ekranda taşınan paragraf görünür, masaüstü tam ekranda panel sayfayı örtmez, UDF → Word telefonda ekrana sığar')

def v112_features(browser,tmp,petition):
    """1.12: themes (picker, palette commands, status bar colour, the paper
    stays white), Word / UYAP line height, previous / next paragraph in the
    panel, toast undo / redo, the touched paragraph glows, drag to move,
    section map, check tour, before / after."""
    ctx=browser.new_context(viewport={'width':1440,'height':900})
    page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    load_page(page,'workspace.html');fetch_data(page,petition)
    # Tema: kartela (6 tema), palet ve mod; kâğıt beyaz kalır; seçim saklanır; durum çubuğu rengi.
    page.locator('#themeToggle').click();page.wait_for_timeout(150)
    assert page.locator('.theme-pop .tp-swatch').count()==6
    page.locator('.tp-swatch[data-palette="bordo"]').click();page.locator('.tp-modes [data-mode="dark"]').click();page.wait_for_timeout(500)
    st=page.evaluate("[document.documentElement.dataset.palette,document.documentElement.dataset.theme,getComputedStyle(document.getElementById('paper')).backgroundColor,localStorage.getItem('wordtoudf.palette'),document.querySelector('meta[name=theme-color]').content]")
    assert st==['bordo','dark','rgb(255, 255, 255)','bordo','#2a0810'],st
    page.keyboard.press('Escape');page.wait_for_timeout(100);assert page.locator('.theme-pop').count()==0
    page.keyboard.press('Control+k');page.keyboard.type('tema grafit');page.wait_for_timeout(100);page.keyboard.press('Enter');page.wait_for_timeout(300)
    assert page.evaluate("document.documentElement.dataset.palette")=='grafit'
    page.evaluate("WordToUDFTheme.setPalette('lacivert');WordToUDFTheme.setMode('auto')");page.wait_for_timeout(450)
    # Satır yüksekliği Word / UYAP gibi: 1,5 satır = 1,5 × 1,15 × 12 pt (20,7 pt = 27,6 px); tek satır 13,8 pt.
    lh=page.evaluate("['explanation-body','metadata'].map(r=>parseFloat(getComputedStyle(document.querySelector(`#paper .paragraph[data-role=\"${r}\"]`)).lineHeight))")
    assert abs(lh[0]-27.6)<.2 and abs(lh[1]-18.4)<.2,lh
    # Panelde önceki / sonraki paragraf (düğme ve Alt+↑/↓).
    page.locator('#paper .paragraph[data-role="explanation-body"]').first.click();page.wait_for_timeout(300)
    title=lambda:page.locator('.edit-sheet .sheet-head small').text_content()
    t0=title();page.locator('.edit-sheet [data-ctl="nav-next"]').click();page.wait_for_timeout(200);t1=title()
    page.keyboard.press('Alt+ArrowUp');page.wait_for_timeout(200)
    assert t1!=t0 and title()==t0,(t0,t1,title())
    page.keyboard.press('Escape');page.wait_for_timeout(100)
    # Seçimle kalın → bildirimde "Geri al", sonra "Yinele"; değişen paragraf parlar.
    assert page.evaluate(SELECT_JS,['#paper .paragraph','verilen karara']);page.wait_for_timeout(300)
    page.locator('.mark-pop .mp-bold').click();page.wait_for_timeout(250)
    bold=lambda:'verilen karara' in page.evaluate("[...document.querySelectorAll('#paper span')].filter(s=>s.style.fontWeight==='700').map(s=>s.textContent)")
    assert bold() and page.locator('#paper .paragraph.glow').count()>=1
    page.locator('.toast:not(.leaving) .t-act',has_text='Geri al').click();page.wait_for_timeout(250);assert not bold()
    page.locator('.toast:not(.leaving) .t-act',has_text='Yinele').click();page.wait_for_timeout(250);assert bold()
    page.keyboard.press('Escape');page.wait_for_timeout(100)
    # Sürükleyerek taşı: tutamaç, bırakma çizgisi, paragraf hedefin altına.
    order=lambda:page.evaluate("[...document.querySelectorAll('#paper .paragraph[data-key]')].map(p=>p.textContent.slice(0,18))")
    before=order()
    page.evaluate("(()=>{const r=document.querySelectorAll('#paper .paragraph[data-key]')[13].getBoundingClientRect();scrollBy(0,r.top-450)})()");page.wait_for_timeout(150)
    page.locator('#paper .paragraph[data-key]').nth(12).hover();page.wait_for_timeout(150)
    g=page.locator('.drag-grip');assert g.is_visible()
    gb=g.bounding_box();db=page.locator('#paper .paragraph[data-key]').nth(15).bounding_box()
    page.mouse.move(gb['x']+10,gb['y']+10);page.mouse.down();page.mouse.move(gb['x']+300,db['y']+db['height']-3,steps=10)
    assert page.locator('.drop-line').is_visible()
    page.mouse.up();page.wait_for_timeout(300)
    after=order();assert after.index(before[12])==15 and sorted(after)==sorted(before),(before[10:18],after[10:18])
    # Bölüm haritası.
    chips=page.locator('#secMap .sec-chip').all_text_contents()
    assert {'Künye','Konu','Açıklamalar','Sonuç ve istem','İmza'}<=set(chips),chips
    # Kontrol turu: madde sayfadaki yerinde işaretlenir; Esc bitirir.
    page.locator('#tourStart').click();page.wait_for_timeout(500)
    assert page.locator('.tour-bar').is_visible() and page.locator('#paper .tour-hit').count()==1
    tb=page.locator('.tour-bar').bounding_box();assert abs(tb['x']+tb['width']/2-720)<2,('tur çubuğu ortada',tb)
    page.keyboard.press('Escape');page.wait_for_timeout(100)
    assert page.locator('.tour-bar').count()==0 and page.locator('#paper .tour-hit').count()==0
    # Önce / sonra: kaynak sayfanın üstünde, ayırıcı klavyeyle de oynar.
    page.locator('#compareToggle').click();page.wait_for_timeout(300)
    assert page.locator('.compare-paper .paragraph').count()>5 and page.locator('.cmp-handle').get_attribute('aria-valuenow')=='50'
    page.locator('.cmp-handle').focus();page.keyboard.press('ArrowLeft');page.keyboard.press('ArrowLeft')
    assert page.locator('.cmp-handle').get_attribute('aria-valuenow')=='40'
    page.locator('#compareToggle').click();assert page.locator('.compare-paper').count()==0
    assert not errors,errors
    ctx.close()
    print('PASS 1.12: 6 tema (kartela, komut, durum çubuğu, beyaz kâğıt), Word/UYAP satır yüksekliği, panelde önceki/sonraki, bildirimde geri al/yinele, parlama, sürükleyerek taşıma, bölüm haritası, kontrol turu, önce/sonra')

def v112_share(browser,petition):
    """1.12: a file shared to the installed web app (Android share sheet)
    reaches the converter through the service worker and a local cache that is
    emptied at once; the real pages with their CSP."""
    import threading,functools,http.server,socketserver,base64
    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self,*a):pass
    srv=socketserver.ThreadingTCPServer(('127.0.0.1',0),functools.partial(Quiet,directory=str(ROOT)));port=srv.server_address[1]
    threading.Thread(target=srv.serve_forever,daemon=True).start()
    try:
        ctx=browser.new_context(viewport={'width':390,'height':844},is_mobile=True,has_touch=True);page=ctx.new_page();errors=[]
        page.on('pageerror',lambda e:errors.append(str(e)))
        page.goto(f'http://127.0.0.1:{port}/workspace.html');page.evaluate("navigator.serviceWorker.ready.then(()=>1)");page.reload();page.wait_for_timeout(500)
        assert page.evaluate("!!navigator.serviceWorker.controller")
        # The share sheet POSTs like a form; a same-origin page without the app's CSP stands in for it here.
        page.goto(f'http://127.0.0.1:{port}/README.md')
        r=page.evaluate("""async b64=>{const bin=Uint8Array.from(atob(b64),c=>c.charCodeAt(0));const fd=new FormData();fd.append('file',new File([bin],'paylasilan.docx'));
          const res=await fetch('share-target',{method:'POST',body:fd,redirect:'manual'});return [res.type,!!(await (await caches.open('wordtoudf-share')).match('shared-file'))];}""",base64.b64encode(Path(petition).read_bytes()).decode())
        assert r==['opaqueredirect',True],r
        page.goto(f'http://127.0.0.1:{port}/workspace.html?shared=1');page.wait_for_timeout(1500)
        assert page.locator('#loadedName').text_content()=='paylasilan.docx' and page.locator('#paper .paragraph').count()>10
        assert page.url.endswith('/workspace.html') and not page.evaluate("caches.open('wordtoudf-share').then(c=>c.match('shared-file')).then(x=>!!x)")
        # Chromium cannot share a .udf from the page: the button stays hidden there.
        assert page.locator('#shareUdf').is_hidden()
        # The check tour bar stays inside a phone screen.
        page.locator('#tourStart').scroll_into_view_if_needed();page.locator('#tourStart').click();page.wait_for_timeout(500)
        tb=page.locator('.tour-bar').bounding_box();assert tb['x']>=0 and tb['x']+tb['width']<=390,tb
        page.locator('.tour-bar .tb-close').click();assert page.locator('.tour-bar').count()==0
        assert not errors,errors
        ctx.close()
        print('PASS 1.12 paylaşım: telefonda paylaşılan Word dosyası uygulamada açılır, geçici kopya hemen silinir')
    finally:
        srv.shutdown()

def csp_check(browser):
    """The other tests strip the pages' CSP; here the real pages are served
    over HTTP with their CSP intact and every security-policy violation or
    console error fails the test."""
    import threading,functools,http.server,socketserver
    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self,*a):pass
    handler=functools.partial(Quiet,directory=str(ROOT))
    srv=socketserver.TCPServer(('127.0.0.1',0),handler);port=srv.server_address[1]
    threading.Thread(target=srv.serve_forever,daemon=True).start()
    try:
        ctx=browser.new_context(accept_downloads=True)
        ctx.add_init_script("window.__csp=[];document.addEventListener('securitypolicyviolation',e=>window.__csp.push(e.violatedDirective+' '+(e.sourceFile||'')+':'+e.lineNumber));")
        page=ctx.new_page();errors=[]
        page.on('pageerror',lambda e:errors.append(str(e)));page.on('console',lambda m:errors.append(m.text) if m.type=='error' else None)
        page.goto(f'http://127.0.0.1:{port}/workspace.html')
        paste_into(page,WORD_HTML,'x');paste_into(page,GDOCS_HTML,'x')
        page.locator('#paper .paragraph').nth(2).click();page.locator('.edit-sheet .chip',has_text='Kalın').click();page.keyboard.press('Escape')
        with page.expect_download(): page.locator('#download').click()
        # PDF: constructed @page stylesheet and the print surface under the real CSP.
        page.evaluate("()=>{window.print=()=>{window.__printSheets=document.adoptedStyleSheets.length;window.dispatchEvent(new Event('afterprint'));};}")
        page.locator('#printPdf').click();page.wait_for_timeout(200)
        assert page.evaluate('window.__printSheets')==1
        v=page.evaluate('window.__csp')
        for name in ['templates.html','reverse.html','index.html','sure.html']:
            page.goto(f'http://127.0.0.1:{port}/{name}');page.wait_for_timeout(300);v+=page.evaluate('window.__csp')
        assert not v,v
        assert not errors,errors
        ctx.close()
        print('PASS güvenlik kuralı (CSP) açık sayfalar: yapıştırma, düzenleyici, indirme, şablon; ihlal ve konsol hatası yok')
    finally:
        srv.shutdown()

if __name__=='__main__':main()
